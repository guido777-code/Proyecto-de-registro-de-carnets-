import streamlit as st
import streamlit.components.v1 as components
import qrcode
import json
import os
import re
import secrets
import base64
import hashlib
import uuid
import textwrap
import html as html_lib
import urllib.parse as urlparse
import cv2
import numpy as np
import pandas as pd
import string
from io import BytesIO
from datetime import datetime, date

APP_VERSION = "5.1"

st.set_page_config(
    page_title="Sistema de Carnet y Asistencia QR",
    page_icon="🪪",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ============================================================
# ESTILO VISUAL V5.1
# ============================================================
st.markdown("""
<style>
    .block-container { padding-top: 1.5rem; padding-bottom: 2rem; }
    [data-testid="stMetric"] {
        border: 1px solid rgba(128,128,128,.22);
        border-radius: 14px;
        padding: 12px 16px;
        background: rgba(128,128,128,.06);
    }
    div.stButton > button, div.stDownloadButton > button {
        border-radius: 10px;
        font-weight: 600;
        min-height: 2.6rem;
    }
    [data-testid="stExpander"] { border-radius: 12px; }
    h1, h2, h3 { letter-spacing: -.3px; }
</style>
""", unsafe_allow_html=True)

# ============================================================
# CONFIGURACIÓN
# ============================================================
OS_DIR = os.getenv("CARNETS_DATA_DIR", "datos_carnets")
os.makedirs(OS_DIR, exist_ok=True)

BASE_URL = os.getenv(
    "CARNETS_BASE_URL",
    "http://localhost:8501"
).rstrip("/")

# Todos los archivos de datos se guardan dentro de la carpeta de datos.
CREDENCIALES_FILE = os.path.join(OS_DIR, "credenciales_admin.json")
DOCENTES_FILE = os.path.join(OS_DIR, "docentes.json")
ASISTENCIA_FILE = os.path.join(OS_DIR, "asistencias.csv")

COLUMNAS_ASISTENCIA = [
    "codigo", "nombres", "apellidos", "nivel", "grado",
    "materia", "docente", "clase_id", "fecha", "hora"
]

MESES_ES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
    5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre",
    12: "Diciembre",
}

NIVELES_EDUCATIVOS = [
    "Preprimaria",
    "Primaria",
    "Básico",
    "Diversificado – Bachillerato",
    "Diversificado – Magisterio",
]

# Catálogo para evitar que los grados/secciones se escriban manualmente
# y reducir errores al crear docentes y asignar estudiantes.
GRADOS_SECCIONES_POR_NIVEL = {
    "Preprimaria": [
        f"Párvulos {n} – Sección {s}"
        for n in range(1, 4) for s in ["A", "B", "C"]
    ],
    "Primaria": [
        f"{nombre} – Sección {s}"
        for nombre in ["Primero", "Segundo", "Tercero", "Cuarto", "Quinto", "Sexto"]
        for s in ["A", "B", "C"]
    ],
    "Básico": [
        f"{nombre} Básico – Sección {s}"
        for nombre in ["Primero", "Segundo", "Tercero"]
        for s in ["A", "B", "C"]
    ],
    "Diversificado – Bachillerato": [
        f"{nombre} Bachillerato – Sección {s}"
        for nombre in ["Cuarto", "Quinto"]
        for s in ["A", "B", "C"]
    ],
    "Diversificado – Magisterio": [
        f"{nombre} Magisterio – Sección {s}"
        for nombre in ["Cuarto", "Quinto", "Sexto"]
        for s in ["A", "B", "C"]
    ],
}

MAX_CLASES_DOCENTE = 12

NOMBRE_COLEGIO_LINEA1 = "COLEGIO MIXTO EVANGÉLICO"
NOMBRE_COLEGIO_LINEA2 = "NAZARENO"
NOMBRE_COLEGIO_LINEA3 = "COBÁN, ALTA VERAPAZ"


# ============================================================
# UTILIDADES GENERALES
# ============================================================
def normalizar_texto(valor):
    return re.sub(r"\s+", " ", str(valor or "").strip())


def normalizar_codigo(codigo):
    return normalizar_texto(codigo).upper()


def normalizar_grado(grado):
    grado = normalizar_texto(grado)
    return grado if grado else "Sin grado"


def ordenar_grados_presentes(grados_presentes):
    extras = sorted(
        {g for g in grados_presentes if g != "Sin grado"},
        key=lambda x: x.casefold()
    )
    return extras + (["Sin grado"] if "Sin grado" in grados_presentes else [])


def codigo_valido(codigo):
    return bool(re.fullmatch(
        r"[A-Za-z0-9_-]{3,30}",
        str(codigo or "").strip()
    ))


def nombre_archivo_seguro(codigo, sufijo, extension):
    codigo_limpio = re.sub(
        r"[^A-Za-z0-9_-]",
        "_",
        str(codigo).strip()
    )
    extension_limpia = re.sub(
        r"[^A-Za-z0-9]",
        "",
        extension.lower()
    )
    return f"{codigo_limpio}_{sufijo}.{extension_limpia}"


def ruta_registro(codigo):
    return os.path.join(OS_DIR, f"{normalizar_codigo(codigo)}.json")


def codigo_existente(codigo):
    return os.path.exists(ruta_registro(codigo))


def get_base64(bin_file):
    with open(bin_file, "rb") as f:
        return base64.b64encode(f.read()).decode()


def escribir_json_atomico(path, datos):
    """Evita dejar un JSON corrupto si el proceso se interrumpe durante la escritura."""
    tmp = f"{path}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


# ============================================================
# CONTRASEÑAS
# ============================================================
def hash_password(password):
    salt = secrets.token_bytes(16)
    iterations = 210_000
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        iterations,
    )
    return f"pbkdf2${iterations}${salt.hex()}${digest.hex()}"


def verificar_hash_password(password, stored_hash):
    if not password or not stored_hash:
        return False

    if stored_hash.startswith("pbkdf2$"):
        try:
            _, iterations, salt_hex, digest_hex = stored_hash.split("$", 3)
            digest = hashlib.pbkdf2_hmac(
                "sha256",
                password.encode("utf-8"),
                bytes.fromhex(salt_hex),
                int(iterations),
            )
            return secrets.compare_digest(
                digest.hex(),
                digest_hex
            )
        except (ValueError, TypeError):
            return False

    # Compatibilidad con instalaciones antiguas SHA-256.
    legacy = hashlib.sha256(
        password.encode("utf-8")
    ).hexdigest()

    return secrets.compare_digest(legacy, stored_hash)


def cargar_credenciales_admin():
    if not os.path.exists(CREDENCIALES_FILE):
        usuario = os.getenv(
            "CARNETS_DEFAULT_USER",
            "admin.nazareno"
        )
        password_inicial = os.getenv("CARNETS_DEFAULT_PASSWORD")

        # Si no existe una contraseña configurada por entorno,
        # se genera una temporal y se muestra UNA sola vez.
        if not password_inicial:
            password_inicial = secrets.token_urlsafe(10)
            st.session_state.password_inicial_temporal = password_inicial

        datos = {
            "usuario": usuario,
            "password_hash": hash_password(password_inicial),
            "requiere_cambio_password": True,
        }

        try:
            escribir_json_atomico(CREDENCIALES_FILE, datos)
        except OSError:
            st.error("No se pudo crear el archivo de credenciales.")
            st.stop()

        return datos

    try:
        with open(CREDENCIALES_FILE, "r", encoding="utf-8") as f:
            datos = json.load(f)

        if (
            not isinstance(datos, dict)
            or "usuario" not in datos
            or "password_hash" not in datos
        ):
            raise ValueError("Formato de credenciales inválido.")

        return datos

    except (json.JSONDecodeError, OSError, ValueError):
        st.error(
            "No se pudo leer el archivo de credenciales. "
            "Revíselo antes de continuar."
        )
        st.stop()


def guardar_credenciales_admin(
    usuario,
    password_hash_value,
    requiere_cambio_password=False
):
    datos = {
        "usuario": usuario,
        "password_hash": password_hash_value,
        "requiere_cambio_password": requiere_cambio_password,
    }
    escribir_json_atomico(CREDENCIALES_FILE, datos)


def verificar_admin(usuario, password):
    credenciales = cargar_credenciales_admin()

    return (
        secrets.compare_digest(
            normalizar_texto(usuario),
            str(credenciales["usuario"])
        )
        and verificar_hash_password(
            password,
            str(credenciales["password_hash"])
        )
    )


# ============================================================
# GESTIÓN DE DOCENTES
# ============================================================
def cargar_docentes():
    if not os.path.exists(DOCENTES_FILE):
        return []

    try:
        with open(DOCENTES_FILE, "r", encoding="utf-8") as f:
            datos = json.load(f)

        return datos if isinstance(datos, list) else []

    except (json.JSONDecodeError, OSError):
        return []


def guardar_docentes(docentes):
    escribir_json_atomico(DOCENTES_FILE, docentes)


def obtener_docente_por_usuario(usuario):
    usuario = normalizar_texto(usuario).casefold()
    for docente in cargar_docentes():
        if normalizar_texto(docente.get("usuario", "")).casefold() == usuario:
            return docente
    return None


def verificar_docente(usuario, password):
    docente = obtener_docente_por_usuario(usuario)
    if not docente:
        return None
    if verificar_hash_password(
        password,
        str(docente.get("password_hash", ""))
    ):
        return docente
    return None


def docente_actual():
    if st.session_state.get("rol") != "docente":
        return None
    return obtener_docente_por_usuario(
        st.session_state.get("username", "")
    )


def parsear_clases_docente(texto):
    clases = []
    vistas = set()

    for linea in str(texto or "").splitlines():
        linea = normalizar_texto(linea)
        if not linea:
            continue

        partes = [normalizar_texto(x) for x in linea.split("|", 1)]
        if len(partes) != 2 or not partes[0] or not partes[1]:
            continue

        grado, materia = partes
        clave = (grado.casefold(), materia.casefold())
        if clave not in vistas:
            clases.append({
                "grado": grado,
                "materia": materia,
            })
            vistas.add(clave)

    return clases


def texto_clases_docente(docente):
    clases = docente.get("clases", [])
    if clases:
        return "\n".join(
            f"{c.get('grado', '')} | {c.get('materia', '')}"
            for c in clases
            if c.get("grado") and c.get("materia")
        )

    # Compatibilidad con docentes creados antes de esta versión.
    materias = docente.get("materias", [])
    grados = docente.get("grados", [])
    if materias and grados:
        return "\n".join(
            f"{grado} | {materia}"
            for grado in grados
            for materia in materias
        )

    return ""


def grado_asignado_a_docente(docente, grado):
    grado = normalizar_grado(grado).casefold()
    for clase in docente.get("clases", []):
        if normalizar_grado(clase.get("grado", "")).casefold() == grado:
            return True
    return False


def materias_asignadas_en_grado(docente, grado):
    grado = normalizar_grado(grado).casefold()
    materias = []

    for clase in docente.get("clases", []):
        if normalizar_grado(clase.get("grado", "")).casefold() == grado:
            materia = normalizar_texto(clase.get("materia", ""))
            if materia and materia not in materias:
                materias.append(materia)

    return materias


# ============================================================
# CLASES / INSCRIPCIONES
# ============================================================
def crear_clase_id(docente_id, grado, materia):
    base = f"{docente_id}|{normalizar_texto(grado).casefold()}|{normalizar_texto(materia).casefold()}"
    return hashlib.sha256(base.encode("utf-8")).hexdigest()[:16]


def inferir_nivel_por_grado(grado):
    g = normalizar_texto(grado).casefold()
    if "párvulo" in g or "parvulo" in g: return "Preprimaria"
    if "básico" in g or "basico" in g: return "Básico"
    if "bachillerato" in g: return "Diversificado – Bachillerato"
    if "magisterio" in g: return "Diversificado – Magisterio"
    if any(x in g for x in ["primero", "segundo", "tercero", "cuarto", "quinto", "sexto"]): return "Primaria"
    return ""


def obtener_clases_docentes(docentes=None):
    docentes = docentes if docentes is not None else cargar_docentes()
    clases = []
    for d in docentes:
        for c in d.get("clases", []):
            grado = normalizar_grado(c.get("grado"))
            materia = normalizar_texto(c.get("materia"))
            if not materia or not d.get("id"):
                continue
            nivel = c.get("nivel") or inferir_nivel_por_grado(grado) or (d.get("niveles", [""])[0] if d.get("niveles") else "")
            clases.append({
                "id": c.get("id") or crear_clase_id(d.get("id"), grado, materia),
                "docente_id": d.get("id"),
                "docente": normalizar_texto(d.get("nombre", "")),
                "nivel": normalizar_texto(nivel),
                "grado": grado,
                "materia": materia,
            })
    return clases


def etiqueta_clase(clase):
    return f"{clase.get('materia')} · {clase.get('grado')} · {clase.get('docente')}"


def obtener_inscripciones_estudiante(registro):
    actuales = registro.get("clases")
    if isinstance(actuales, list) and actuales:
        salida = []
        for c in actuales:
            if not isinstance(c, dict):
                continue
            grado = normalizar_grado(c.get("grado"))
            materia = normalizar_texto(c.get("materia"))
            docente = normalizar_texto(c.get("docente"))
            nivel = normalizar_texto(c.get("nivel"))
            if grado and materia and materia != "No asignada":
                salida.append({
                    "id": c.get("id") or crear_clase_id(c.get("docente_id", ""), grado, materia),
                    "docente_id": c.get("docente_id", ""),
                    "docente": docente,
                    "nivel": nivel,
                    "grado": grado,
                    "materia": materia,
                })
        if salida:
            return salida

    # Compatibilidad con carnets creados en versiones anteriores.
    materia = normalizar_texto(registro.get("materia", ""))
    docente = normalizar_texto(registro.get("docente", ""))
    grado = normalizar_grado(registro.get("grado", ""))
    nivel = normalizar_texto(registro.get("nivel", ""))
    if materia and materia != "No asignada":
        docente_id = ""
        for d in cargar_docentes():
            if normalizar_texto(d.get("nombre", "")).casefold() == docente.casefold():
                docente_id = d.get("id", "")
                break
        return [{
            "id": crear_clase_id(docente_id, grado, materia),
            "docente_id": docente_id,
            "docente": docente,
            "nivel": nivel,
            "grado": grado,
            "materia": materia,
        }]
    return []


def estudiante_en_clase(registro, clase):
    cid = clase.get("id")
    return any(
        c.get("id") == cid
        or (normalizar_texto(c.get("docente")).casefold() == normalizar_texto(clase.get("docente")).casefold()
            and normalizar_texto(c.get("grado")).casefold() == normalizar_texto(clase.get("grado")).casefold()
            and normalizar_texto(c.get("materia")).casefold() == normalizar_texto(clase.get("materia")).casefold())
        for c in obtener_inscripciones_estudiante(registro)
    )


def clases_a_texto(clases):
    return "\n".join(
        f"{c.get('nivel', '')} | {c.get('grado', '')} | {c.get('materia', '')}"
        for c in clases
    )


# ============================================================
# CARNETS
# ============================================================
def cargar_registro(codigo):
    codigo = normalizar_codigo(codigo)

    if not codigo_valido(codigo):
        return None

    path = ruta_registro(codigo)

    if not os.path.exists(path):
        return None

    try:
        with open(path, "r", encoding="utf-8") as f:
            datos = json.load(f)

        return datos if isinstance(datos, dict) else None

    except (json.JSONDecodeError, OSError):
        return None


def listar_registros():
    registros = []

    try:
        archivos = sorted(os.listdir(OS_DIR))
    except OSError:
        return registros

    for fname in archivos:
        if not fname.lower().endswith(".json"):
            continue

        # Ignorar archivos internos.
        if fname in {
            "credenciales_admin.json",
            "docentes.json",
        }:
            continue

        path = os.path.join(OS_DIR, fname)

        try:
            with open(path, "r", encoding="utf-8") as f:
                datos = json.load(f)

            if isinstance(datos, dict) and datos.get("codigo"):
                registros.append(datos)

        except (json.JSONDecodeError, OSError):
            continue

    return registros


def buscar_duplicado_por_dpi(dpi):
    dpi_normalizado = normalizar_texto(dpi).lower().replace(" ", "")

    if not dpi_normalizado or dpi_normalizado == "noaplica":
        return None

    for registro in listar_registros():
        dpi_existente = (
            normalizar_texto(registro.get("dpi"))
            .lower()
            .replace(" ", "")
        )

        if (
            dpi_existente
            and dpi_existente != "noaplica"
            and dpi_existente == dpi_normalizado
        ):
            return registro

    return None


def generar_codigo_aleatorio(longitud=7):
    caracteres = string.ascii_uppercase + string.digits

    for _ in range(1000):
        codigo = "".join(
            secrets.choice(caracteres)
            for _ in range(longitud)
        )

        if not codigo_existente(codigo):
            return codigo

    raise RuntimeError(
        "No fue posible generar un código único."
    )


# ============================================================
# ASISTENCIA
# ============================================================
def cargar_asistencias():
    if not os.path.exists(ASISTENCIA_FILE):
        return pd.DataFrame(columns=COLUMNAS_ASISTENCIA)

    try:
        df = pd.read_csv(
            ASISTENCIA_FILE,
            dtype=str,
            encoding="utf-8-sig"
        )

    except (pd.errors.EmptyDataError, FileNotFoundError):
        return pd.DataFrame(columns=COLUMNAS_ASISTENCIA)

    for col in COLUMNAS_ASISTENCIA:
        if col not in df.columns:
            df[col] = ""

    return df[COLUMNAS_ASISTENCIA].fillna("")


def guardar_asistencias(df):
    salida = df.copy()

    for col in COLUMNAS_ASISTENCIA:
        if col not in salida.columns:
            salida[col] = ""

    salida = salida[COLUMNAS_ASISTENCIA].fillna("").astype(str)

    tmp = f"{ASISTENCIA_FILE}.tmp"
    salida.to_csv(
        tmp,
        index=False,
        encoding="utf-8-sig"
    )
    os.replace(tmp, ASISTENCIA_FILE)


def registrar_asistencia_codigo(codigo, clase=None):
    codigo = normalizar_codigo(codigo)
    datos = cargar_registro(codigo)
    if not datos:
        return False, "Código no registrado en la base de datos."
    if not clase:
        return False, "Seleccione una clase antes de registrar la asistencia."
    if not estudiante_en_clase(datos, clase):
        return False, (
            f"El estudiante no está inscrito en {clase.get('materia')} "
            f"· {clase.get('grado')}."
        )

    df = cargar_asistencias()
    ahora = datetime.now()
    fecha_hoy = ahora.strftime("%Y-%m-%d")
    hora_actual = ahora.strftime("%H:%M:%S")
    clase_id = clase.get("id") or crear_clase_id(
        clase.get("docente_id", ""), clase.get("grado", ""), clase.get("materia", "")
    )

    # Un estudiante puede asistir a varias materias el mismo día,
    # pero no dos veces a la misma clase en la misma fecha.
    if not df.empty:
        ya = not df[
            (df["codigo"].astype(str).str.upper() == codigo)
            & (df["fecha"].astype(str) == fecha_hoy)
            & (df["clase_id"].astype(str) == str(clase_id))
        ].empty
    else:
        ya = False

    nombre = normalizar_texto(f"{datos.get('nombres', '')} {datos.get('apellidos', '')}")
    if ya:
        return False, (
            f"{nombre} ya registró asistencia hoy en "
            f"{clase.get('materia')} · {clase.get('grado')}."
        )

    nueva_fila = {
        "codigo": codigo,
        "nombres": datos.get("nombres", ""),
        "apellidos": datos.get("apellidos", ""),
        "nivel": clase.get("nivel", datos.get("nivel", "")),
        "grado": clase.get("grado", ""),
        "materia": clase.get("materia", ""),
        "docente": clase.get("docente", ""),
        "clase_id": clase_id,
        "fecha": fecha_hoy,
        "hora": hora_actual,
    }
    df = pd.concat([df, pd.DataFrame([nueva_fila])], ignore_index=True)
    guardar_asistencias(df)
    return True, (
        f"Asistencia registrada: {nombre} · {clase.get('materia')} · "
        f"{clase.get('grado')}"
    )


def generar_excel_bytes(
    df,
    titulo_reporte="HISTORIAL DE ASISTENCIA"
):
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.drawing.image import Image
        from openpyxl.utils import get_column_letter

    except ImportError:
        st.error(
            "Falta la librería openpyxl. "
            "Instálela con: pip install openpyxl"
        )
        return None

    df = df.copy()

    for col in COLUMNAS_ASISTENCIA:
        if col not in df.columns:
            df[col] = ""

    df = df[COLUMNAS_ASISTENCIA].fillna("")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial Asistencia"
    ws.freeze_panes = "B8"
    ws.sheet_view.showGridLines = True

    # Logo
    if LOGO_PATH and os.path.exists(LOGO_PATH):
        try:
            img = Image(LOGO_PATH)
            img.width = 85
            img.height = 85
            ws.add_image(img, "B2")
        except Exception:
            pass

    ws["D2"] = "COLEGIO MIXTO EVANGÉLICO NAZARENO"
    ws["D2"].font = Font(
        name="Calibri",
        size=15,
        bold=True,
        color="002060"
    )

    ws["D3"] = "COBÁN, ALTA VERAPAZ"
    ws["D3"].font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color="595959"
    )

    ws["D4"] = titulo_reporte.upper()
    ws["D4"].font = Font(
        name="Calibri",
        size=12,
        bold=True,
        color="1F4E79"
    )

    fila_encabezado = 7
    columna_inicio = 2

    fill_header = PatternFill(
        start_color="1F4E79",
        end_color="1F4E79",
        fill_type="solid"
    )

    font_header = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF"
    )

    fill_zebra = PatternFill(
        start_color="F2F5F9",
        end_color="F2F5F9",
        fill_type="solid"
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers = [
        "Código", "Nombres", "Apellidos", "Nivel", "Grado",
        "Materia", "Docente", "Clase ID", "Fecha", "Hora",
    ]

    for idx, text in enumerate(headers):
        col = columna_inicio + idx
        cell = ws.cell(
            row=fila_encabezado,
            column=col,
            value=text
        )
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.border = thin_border

    for row_idx, row_data in enumerate(
        df.to_dict(orient="records")
    ):
        fila_actual = fila_encabezado + 1 + row_idx

        valores = [
            row_data.get("codigo", ""), row_data.get("nombres", ""),
            row_data.get("apellidos", ""), row_data.get("nivel", ""),
            row_data.get("grado", ""), row_data.get("materia", ""),
            row_data.get("docente", ""), row_data.get("clase_id", ""),
            row_data.get("fecha", ""), row_data.get("hora", ""),
        ]

        for col_idx, valor in enumerate(valores):
            col = columna_inicio + col_idx

            cell = ws.cell(
                row=fila_actual,
                column=col,
                value=str(valor)
            )

            cell.font = Font(
                name="Calibri",
                size=10
            )
            cell.border = thin_border

            if row_idx % 2 == 1:
                cell.fill = fill_zebra

            cell.alignment = Alignment(
                horizontal=(
                    "center"
                    if col_idx in [0, 3, 7, 8, 9]
                    else "left"
                ),
                vertical="center"
            )

    for col_idx in range(
        columna_inicio,
        columna_inicio + len(headers)
    ):
        col_letter = get_column_letter(col_idx)
        max_len = 0

        for row in range(
            fila_encabezado,
            ws.max_row + 1
        ):
            value = ws.cell(
                row=row,
                column=col_idx
            ).value

            if value is not None:
                max_len = max(
                    max_len,
                    len(str(value))
                )

        ws.column_dimensions[col_letter].width = max(
            min(max_len + 4, 45),
            12
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# FOTOS Y CARNET
# ============================================================
def obtener_foto_base64(nombre_foto):
    if not nombre_foto:
        return None

    ruta = os.path.join(OS_DIR, os.path.basename(nombre_foto))

    if not os.path.exists(ruta):
        return None

    extension = (
        os.path.splitext(nombre_foto)[1]
        .lower()
        .replace(".", "")
    )

    mime = {
        "jpg": "jpeg",
        "jpeg": "jpeg",
        "png": "png",
    }.get(extension)

    if not mime:
        return None

    try:
        with open(ruta, "rb") as f:
            b64 = base64.b64encode(f.read()).decode()

        return f"data:image/{mime};base64,{b64}"

    except OSError:
        return None


def renderizar_carnet(datos):
    nombres = html_lib.escape(
        str(datos.get("nombres", ""))
    )
    apellidos = html_lib.escape(
        str(datos.get("apellidos", ""))
    )
    codigo = html_lib.escape(
        str(datos.get("codigo", ""))
    )
    grado = html_lib.escape(
        str(datos.get("grado", ""))
    )

    foto_b64 = obtener_foto_base64(
        datos.get("foto")
    )

    qr_path = os.path.join(
        OS_DIR,
        f"{normalizar_codigo(datos.get('codigo', ''))}_qr.png"
    )

    qr_b64 = (
        get_base64(qr_path)
        if os.path.exists(qr_path)
        else ""
    )

    logo_b64 = ""

    if LOGO_PATH and os.path.exists(LOGO_PATH):
        ext_logo = (
            os.path.splitext(LOGO_PATH)[1]
            .lower()
            .replace(".", "")
        )

        mime_logo = (
            "jpeg"
            if ext_logo in {"jpg", "jpeg"}
            else "png"
        )

        logo_b64 = (
            f"data:image/{mime_logo};base64,"
            f"{get_base64(LOGO_PATH)}"
        )

    if foto_b64:
        img_foto = (
            '<img src="'
            f'{foto_b64}'
            '" style="width:130px;height:150px;'
            'object-fit:cover;border-radius:12px;'
            'border:3px solid #fff;" />'
        )
    else:
        img_foto = (
            '<div style="width:130px;height:150px;'
            'border-radius:12px;background:#e2e8f0;'
            'display:flex;align-items:center;'
            'justify-content:center;font-size:40px;'
            'border:3px solid #fff;">👤</div>'
        )

    img_qr = (
        f'<img src="data:image/png;base64,{qr_b64}" '
        'style="width:65px;height:65px;border-radius:6px;'
        'background:white;padding:2px;" />'
        if qr_b64
        else ""
    )

    logo_img = (
        f'<img src="{logo_b64}" '
        'style="width:40px;height:40px;'
        'border-radius:50%;background:white;padding:2px;" />'
        if logo_b64
        else ""
    )

    html_content = textwrap.dedent(
        f"""
        <!DOCTYPE html>
        <html>
        <head>
            <style>
                body {{
                    margin: 0;
                    padding: 0;
                    background: transparent;
                    font-family: sans-serif;
                    display: flex;
                    justify-content: center;
                }}

                .card {{
                    width: 300px;
                    border-radius: 20px;
                    background:
                        linear-gradient(
                            160deg,
                            #0c2340 0%,
                            #0066ff 100%
                        );
                    padding: 18px;
                    color: #ffffff;
                    box-shadow:
                        0 8px 20px rgba(0,0,0,0.3);
                    box-sizing: border-box;
                }}

                .header {{
                    display: flex;
                    align-items: center;
                    gap: 10px;
                    margin-bottom: 8px;
                }}

                .title {{
                    text-align: left;
                    font-size: 11px;
                    font-weight: bold;
                    line-height: 1.2;
                }}

                .subtitle {{
                    text-align: left;
                    margin-bottom: 10px;
                }}

                .sub-label {{
                    font-size: 9px;
                    opacity: 0.8;
                }}

                .sub-year {{
                    font-size: 15px;
                    font-weight: bold;
                }}

                .photo-container {{
                    display: flex;
                    justify-content: center;
                    margin-bottom: 10px;
                }}

                .footer {{
                    display: flex;
                    justify-content: space-between;
                    align-items: flex-end;
                    text-align: left;
                    gap: 8px;
                }}

                .field-label {{
                    font-size: 8px;
                    opacity: 0.8;
                }}

                .field-val {{
                    font-size: 13px;
                    font-weight: bold;
                    margin-bottom: 4px;
                }}
            </style>
        </head>

        <body>
            <div class="card">
                <div class="header">
                    {logo_img}
                    <div class="title">
                        Colegio Mixto<br>
                        Evangélico Nazareno
                    </div>
                </div>

                <div class="subtitle">
                    <div class="sub-label">
                        Carnet estudiantil
                    </div>
                    <div class="sub-year">
                        Ciclo escolar {datetime.now().year}
                    </div>
                </div>

                <div class="photo-container">
                    {img_foto}
                </div>

                <div class="footer">
                    <div>
                        <div class="field-label">CÓDIGO</div>
                        <div class="field-val">
                            {codigo}
                        </div>

                        <div class="field-label">
                            NOMBRE
                        </div>
                        <div class="field-val">
                            {nombres}<br>{apellidos}
                        </div>

                        <div class="field-label">
                            GRADO
                        </div>
                        <div
                            class="field-val"
                            style="margin-bottom:0;"
                        >
                            {grado}
                        </div>
                    </div>

                    <div>
                        {img_qr}
                    </div>
                </div>
            </div>
        </body>
        </html>
        """
    )

    components.html(
        html_content,
        height=440
    )


def mostrar_tabla_con_carnet(
    df_tabla,
    key_tabla,
    permitir_edicion=False
):
    if df_tabla.empty:
        st.info("No hay registros para mostrar.")
        return

    df_tabla = df_tabla.reset_index(drop=True)

    columnas_visibles = [
        c for c in [
            "codigo", "nombres", "apellidos", "nivel", "grado",
            "materia", "docente", "fecha", "hora",
        ]
        if c in df_tabla.columns
    ]

    df_visual = df_tabla[columnas_visibles]

    evento = st.dataframe(
        df_visual,
        use_container_width=True,
        hide_index=True,
        on_select="rerun",
        selection_mode="single-row",
        key=key_tabla,
    )

    filas = []

    if evento is not None:
        try:
            filas = list(
                evento.selection.rows
            )
        except AttributeError:
            try:
                filas = evento.selection.get(
                    "rows",
                    []
                )
            except AttributeError:
                filas = []

    if filas and "codigo" in df_visual.columns:
        codigo_tocado = normalizar_codigo(
            df_visual.iloc[filas[0]]["codigo"]
        )

        datos_tocados = cargar_registro(
            codigo_tocado
        )

        if datos_tocados:
            st.session_state[f"codigo_seleccionado_{key_tabla}"] = (
                codigo_tocado
            )

            st.write("")

            c1, c2, c3 = st.columns(3)

            with c2:
                renderizar_carnet(
                    datos_tocados
                )

    if permitir_edicion:
        codigo_seleccionado = st.session_state.get(
            f"codigo_seleccionado_{key_tabla}"
        )

        if codigo_seleccionado:
            datos_seleccionados = cargar_registro(
                codigo_seleccionado
            )

            if datos_seleccionados:
                st.markdown("---")

                if st.button(
                    "✏️ Editar carnet seleccionado",
                    type="primary",
                    use_container_width=True,
                    key=f"btn_editar_carnet_{key_tabla}"
                ):
                    st.session_state.carnet_editando = (
                        codigo_seleccionado
                    )
                    st.rerun()


# ============================================================
# LOGO / FONDO
# ============================================================
LOGO_PATH = os.getenv(
    "CARNETS_LOGO_PATH",
    "logo_colegio_nazareno.png"
)

if not os.path.exists(LOGO_PATH):
    LOGO_PATH = None

    for ext in [
        "logo.png",
        "logo.jpg",
        "logo.jpeg"
    ]:
        if os.path.exists(ext):
            LOGO_PATH = ext
            break

FONDO_PATH = None

for ext in [
    "fondo.jpg",
    "fondo.jpeg",
    "fondo.png"
]:
    if os.path.exists(ext):
        FONDO_PATH = ext
        break


# ============================================================
# ESTADO DE SESIÓN
# ============================================================
defaults = {
    "pagina": "login_admin",
    "logged_in": False,
    "rol": None,
    "username": None,
    "docente_id": None,
    "ultimo_codigo_registrado": None,
    "codigo_escaneado": None,
    "modo_oscuro": False,
    "docente_editando": None,
    "carnet_editando": None,
    "vista_actual": "inicio",
    "password_inicial_temporal": None,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# TEMA
# ============================================================
MODO_OSCURO = st.session_state.modo_oscuro

if MODO_OSCURO:
    PALETA = {
        "texto": "#F4F7FB",
        "texto_suave": "#A8B3C7",
        "fondo_gradiente": (
            "radial-gradient(circle at 15% 0%, rgba(47,111,237,.16), transparent 30%), "
            "linear-gradient(135deg,#07111F 0%,#0B172A 48%,#101E35 100%)"
        ),
        "input_bg": "#111F33",
        "input_borde": "#30445F",
        "placeholder": "#71809A",
        "primario": "#4F8CFF",
        "primario_texto": "#FFFFFF",
        "secundario_bg": "#132238",
        "card_bg": "rgba(15,29,49,.94)",
        "card_borde": "rgba(110,145,190,.20)",
        "sidebar_bg": "#071426",
        "sidebar_texto": "#D8E2F0",
        "sidebar_activo_bg": "#2F6FED",
        "sidebar_hover_bg": "#102744",
        "badge_bg": "#12233B",
        "hr_color": "rgba(125,151,184,.18)",
    }
else:
    PALETA = {
        "texto": "#102A43",
        "texto_suave": "#62748A",
        "fondo_gradiente": (
            "radial-gradient(circle at 12% 0%, rgba(47,111,237,.12), transparent 30%), "
            "linear-gradient(135deg,#F5F9FE 0%,#EEF5FC 48%,#E5F0FA 100%)"
        ),
        "input_bg": "#FFFFFF",
        "input_borde": "#D8E3EF",
        "placeholder": "#9AA9BB",
        "primario": "#246BCE",
        "primario_texto": "#FFFFFF",
        "secundario_bg": "#FFFFFF",
        "card_bg": "rgba(255,255,255,.94)",
        "card_borde": "rgba(31,73,125,.10)",
        "sidebar_bg": "#0B2A50",
        "sidebar_texto": "#E5EEF8",
        "sidebar_activo_bg": "#2F6FED",
        "sidebar_hover_bg": "#143B69",
        "badge_bg": "#F1F6FC",
        "hr_color": "rgba(31,73,125,.10)",
    }

st.markdown(
    f"""
    <style>
    /* ========================================================
       IDENTIDAD VISUAL PROFESIONAL
       Solo estilos: no modifica la lógica de la aplicación.
       ======================================================== */

    :root {{
        --primary: {PALETA["primario"]};
        --text: {PALETA["texto"]};
        --muted: {PALETA["texto_suave"]};
        --surface: {PALETA["card_bg"]};
        --surface-2: {PALETA["input_bg"]};
        --border: {PALETA["card_borde"]};
        --input-border: {PALETA["input_borde"]};
        --sidebar: {PALETA["sidebar_bg"]};
        --sidebar-text: {PALETA["sidebar_texto"]};
        --sidebar-active: {PALETA["sidebar_activo_bg"]};
        --sidebar-hover: {PALETA["sidebar_hover_bg"]};
        --shadow-sm: 0 4px 14px rgba(15,42,82,.06);
        --shadow-md: 0 14px 38px rgba(15,42,82,.10);
        --shadow-lg: 0 22px 60px rgba(15,42,82,.14);
        --radius: 18px;
    }}

    * {{
        box-sizing: border-box;
    }}

    html, body, [data-testid="stAppViewContainer"] {{
        background: {PALETA["fondo_gradiente"]} !important;
    }}

    [data-testid="stAppViewContainer"] {{
        min-height: 100vh;
    }}

    [data-testid="stHeader"] {{
        background: transparent !important;
        height: 0 !important;
    }}

    [data-testid="stToolbar"] {{
        visibility: hidden;
        height: 0;
    }}

    .block-container {{
        max-width: 1320px !important;
        padding: 1.35rem 2rem 2.8rem !important;
    }}

    h1, h2, h3, h4, h5, h6, label, p,
    [data-testid="stMarkdownContainer"] p,
    [data-testid="stMarkdownContainer"] li,
    [data-testid="stCaptionContainer"] {{
        color: {PALETA["texto"]} !important;
        font-family: "Segoe UI", Inter, system-ui, -apple-system, BlinkMacSystemFont, sans-serif !important;
    }}

    h1, h2, h3, h4, h5, h6 {{
        letter-spacing: -.02em;
    }}

    h1 {{ font-weight: 800 !important; }}
    h2, h3 {{ font-weight: 750 !important; }}

    hr {{
        border: 0 !important;
        border-top: 1px solid {PALETA["hr_color"]} !important;
        margin: 1.15rem 0 !important;
    }}

    /* Tarjetas */
    .card-panel,
    .login-card {{
        position: relative;
        background: {PALETA["card_bg"]} !important;
        border: 1px solid {PALETA["card_borde"]} !important;
        border-radius: var(--radius) !important;
        box-shadow: var(--shadow-md) !important;
        backdrop-filter: blur(16px);
        -webkit-backdrop-filter: blur(16px);
    }}

    .card-panel {{
        padding: 28px 30px !important;
        margin-bottom: 20px !important;
    }}

    .login-card {{
        padding: 42px !important;
        box-shadow: var(--shadow-lg) !important;
    }}

    .card-panel::before,
    .login-card::before {{
        content: "";
        position: absolute;
        top: 0;
        left: 24px;
        right: 24px;
        height: 2px;
        border-radius: 999px;
        background: linear-gradient(90deg, transparent, var(--primary), transparent);
        opacity: .55;
        pointer-events: none;
    }}

    /* Inputs */
    div[data-baseweb="input"],
    div[data-baseweb="base-input"],
    div[data-baseweb="select"] > div,
    [data-testid="stTextInput"] input,
    [data-testid="stDateInput"] input,
    input[type="text"],
    input[type="password"],
    input[type="date"] {{
        background: {PALETA["input_bg"]} !important;
        color: {PALETA["texto"]} !important;
        border-color: {PALETA["input_borde"]} !important;
        border-radius: 12px !important;
        min-height: 44px;
        transition: border-color .18s ease, box-shadow .18s ease, transform .18s ease;
    }}

    div[data-baseweb="input"]:focus-within,
    div[data-baseweb="base-input"]:focus-within,
    [data-testid="stTextInput"]:focus-within,
    [data-testid="stDateInput"]:focus-within {{
        border-color: var(--primary) !important;
        box-shadow: 0 0 0 3px rgba(47,111,237,.12) !important;
    }}

    input::placeholder {{
        color: {PALETA["placeholder"]} !important;
    }}

    div[data-baseweb="select"] svg,
    div[data-baseweb="datepicker"] svg,
    [data-testid="stDateInput"] svg,
    [data-testid="stSelectbox"] svg {{
        fill: {PALETA["texto"]} !important;
    }}

    /* Selectores / popovers */
    div[data-baseweb="popover"] ul[role="listbox"],
    div[data-baseweb="popover"] div[data-baseweb="calendar"],
    div[data-baseweb="menu"] {{
        background: {PALETA["input_bg"]} !important;
        border: 1px solid {PALETA["input_borde"]} !important;
        border-radius: 12px !important;
        box-shadow: var(--shadow-lg) !important;
    }}

    div[data-baseweb="popover"] ul[role="listbox"] *,
    div[data-baseweb="popover"] div[data-baseweb="calendar"] *,
    div[data-baseweb="menu"] * {{
        color: {PALETA["texto"]} !important;
    }}

    div[data-baseweb="popover"] li[aria-selected="true"],
    div[data-baseweb="popover"] div[aria-selected="true"] {{
        background: rgba(47,111,237,.12) !important;
    }}

    /* Botones */
    div.stButton > button,
    div.stDownloadButton > button,
    div.stFormSubmitButton > button {{
        min-height: 44px !important;
        border-radius: 12px !important;
        padding: 0 18px !important;
        font-weight: 700 !important;
        letter-spacing: .01em;
        transition: transform .16s ease, box-shadow .16s ease, border-color .16s ease !important;
    }}

    div.stButton > button:hover,
    div.stDownloadButton > button:hover,
    div.stFormSubmitButton > button:hover {{
        transform: translateY(-1px);
        box-shadow: 0 8px 18px rgba(15,42,82,.12) !important;
    }}

    div.stButton > button[kind="primary"],
    div.stDownloadButton > button,
    div.stFormSubmitButton > button[kind="primary"] {{
        background: linear-gradient(135deg, {PALETA["primario"]}, #4F8CFF) !important;
        color: #FFFFFF !important;
        border: 0 !important;
        box-shadow: 0 7px 18px rgba(47,111,237,.22) !important;
    }}

    div.stButton > button[kind="secondary"],
    div.stFormSubmitButton > button[kind="secondary"] {{
        background: {PALETA["secundario_bg"]} !important;
        color: {PALETA["texto"]} !important;
        border: 1px solid {PALETA["input_borde"]} !important;
    }}

    /* Métricas */
    [data-testid="stMetric"] {{
        background: linear-gradient(145deg, rgba(255,255,255,.05), rgba(255,255,255,0));
        border: 1px solid {PALETA["card_borde"]};
        border-radius: 14px;
        padding: 16px 18px;
        box-shadow: var(--shadow-sm);
    }}

    [data-testid="stMetricLabel"] {{
        color: {PALETA["texto_suave"]} !important;
        font-weight: 600 !important;
    }}

    [data-testid="stMetricValue"] {{
        color: {PALETA["texto"]} !important;
        font-weight: 800 !important;
    }}

    /* Alertas */
    [data-testid="stAlert"] {{
        border-radius: 13px !important;
        border: 1px solid {PALETA["card_borde"]} !important;
        box-shadow: var(--shadow-sm) !important;
    }}

    /* Expander */
    [data-testid="stExpander"] {{
        background: {PALETA["card_bg"]} !important;
        border: 1px solid {PALETA["card_borde"]} !important;
        border-radius: 14px !important;
        overflow: hidden;
        box-shadow: var(--shadow-sm) !important;
    }}

    [data-testid="stExpander"] summary,
    [data-testid="stExpander"] details > summary {{
        background: transparent !important;
    }}

    [data-testid="stExpander"] summary,
    [data-testid="stExpander"] summary * {{
        color: {PALETA["texto"]} !important;
        font-weight: 700 !important;
    }}

    /* Sidebar */
    [data-testid="stSidebar"],
    [data-testid="stSidebarContent"] {{
        background: linear-gradient(180deg, {PALETA["sidebar_bg"]} 0%, #0D3764 100%) !important;
        border-right: 1px solid rgba(255,255,255,.06) !important;
    }}

    [data-testid="stSidebar"] * {{
        color: {PALETA["sidebar_texto"]} !important;
    }}

    [data-testid="stSidebar"] hr {{
        border-top: 1px solid rgba(255,255,255,.11) !important;
    }}

    [data-testid="stSidebar"] div.stButton > button {{
        width: 100%;
        min-height: 45px !important;
        background: transparent !important;
        border: 1px solid transparent !important;
        color: #FFFFFF !important;
        text-align: left !important;
        justify-content: flex-start !important;
        opacity: 1 !important;
        border-radius: 12px !important;
        font-weight: 600 !important;
        padding: 10px 14px !important;
        margin: 2px 0 !important;
        box-shadow: none !important;
    }}

    /* Refuerzo de contraste: evita que Streamlit aplique texto oscuro a los botones secundarios */
    [data-testid="stSidebar"] div.stButton > button,
    [data-testid="stSidebar"] div.stButton > button *,
    [data-testid="stSidebar"] div.stButton > button[kind="secondary"],
    [data-testid="stSidebar"] div.stButton > button[kind="secondary"] * {{
        color: #FFFFFF !important;
        opacity: 1 !important;
    }}

    [data-testid="stSidebar"] div.stButton > button:hover {{
        background: {PALETA["sidebar_hover_bg"]} !important;
        border-color: rgba(255,255,255,.06) !important;
        transform: translateX(2px);
        box-shadow: none !important;
    }}

    [data-testid="stSidebar"] div.stButton > button[kind="primary"] {{
        background: linear-gradient(135deg, {PALETA["sidebar_activo_bg"]}, #4F8CFF) !important;
        color: #FFFFFF !important;
        box-shadow: 0 8px 20px rgba(47,111,237,.22) !important;
    }}

    [data-testid="stSidebar"] div.stButton > button p {{
        text-align: left !important;
        font-weight: 650 !important;
    }}

    .sidebar-nombre-colegio {{
        line-height: 1.2;
        margin: 8px 0 14px;
        padding: 0 4px;
    }}

    .sidebar-nombre-colegio .l1 {{
        font-size: 10px;
        font-weight: 700;
        letter-spacing: .08em;
        opacity: .72;
    }}

    .sidebar-nombre-colegio .l2 {{
        font-size: 22px;
        font-weight: 850;
        letter-spacing: .03em;
    }}

    .sidebar-nombre-colegio .l3 {{
        font-size: 9px;
        opacity: .62;
        letter-spacing: .13em;
        margin-top: 3px;
    }}

    /* Usuario y barra superior */
    .badge-usuario {{
        background: {PALETA["badge_bg"]} !important;
        color: {PALETA["texto"]} !important;
        border: 1px solid {PALETA["card_borde"]} !important;
        border-radius: 999px !important;
        padding: 8px 15px !important;
        font-size: 13px;
        font-weight: 700;
        display: inline-flex;
        align-items: center;
        gap: 6px;
        white-space: nowrap;
        box-shadow: var(--shadow-sm);
    }}

    .badge-usuario span {{
        color: {PALETA["texto"]} !important;
    }}

    /* Cámara y carga de archivos */
    [data-testid="stCameraInput"],
    [data-testid="stFileUploader"] section {{
        background: {PALETA["input_bg"]} !important;
        border: 1px dashed {PALETA["input_borde"]} !important;
        border-radius: 13px !important;
    }}

    [data-testid="stCameraInput"] p,
    [data-testid="stCameraInput"] span,
    [data-testid="stCameraInput"] label,
    [data-testid="stFileUploader"] p,
    [data-testid="stFileUploader"] span {{
        color: {PALETA["texto"]} !important;
    }}

    /* Dataframes */
    [data-testid="stDataFrame"] {{
        border: 1px solid {PALETA["card_borde"]} !important;
        border-radius: 14px !important;
        overflow: hidden !important;
        box-shadow: var(--shadow-sm);
    }}

    /* Tabs */
    [data-baseweb="tab-list"] {{
        gap: 6px !important;
        border-bottom: 1px solid {PALETA["hr_color"]} !important;
    }}

    [data-baseweb="tab-list"] button[data-baseweb="tab"],
    [data-baseweb="tab-list"] [role="tab"] {{
        min-height: 46px !important;
        padding: 10px 16px !important;
        border-radius: 10px 10px 0 0 !important;
        color: {PALETA["texto_suave"]} !important;
        font-weight: 650 !important;
    }}

    [data-baseweb="tab-list"] button[data-baseweb="tab"][aria-selected="true"],
    [data-baseweb="tab-list"] [role="tab"][aria-selected="true"] {{
        color: var(--primary) !important;
        background: rgba(47,111,237,.07) !important;
    }}

    /* Responsive */
    @media (max-width: 900px) {{
        .block-container {{
            padding: 1rem 1rem 2rem !important;
        }}
        .card-panel, .login-card {{
            padding: 22px !important;
        }}
        .sidebar-nombre-colegio .l2 {{
            font-size: 19px;
        }}
    }}
    </style>
    """,
    unsafe_allow_html=True
)


# BARRA SUPERIOR
# ============================================================
def barra_superior():
    _, col_der = st.columns([6, 1.6])

    with col_der:
        c1, c2 = st.columns([3.4, 1], gap="small")

        with c1:
            st.markdown(
                f"""
                <div class='badge-usuario'>
                    👤 Usuario: {html_lib.escape(
                        str(st.session_state.username or "")
                    )}
                </div>
                """,
                unsafe_allow_html=True,
            )

        with c2:
            icono_tema = (
                "☀️"
                if MODO_OSCURO
                else "🌙"
            )

            if st.button(
                icono_tema,
                key="btn_toggle_tema",
                help="Cambiar tema"
            ):
                st.session_state.modo_oscuro = (
                    not st.session_state.modo_oscuro
                )
                st.rerun()

    st.write("")


# ============================================================
# LOGIN
# ============================================================
def pantalla_login_admin():
    col_a, col_b, col_c = st.columns([1, 1.2, 1])

    with col_b:
        st.markdown(
            "<div class='login-card'>",
            unsafe_allow_html=True
        )

        if LOGO_PATH:
            lc1, lc2, lc3 = st.columns(3)
            with lc2:
                st.image(LOGO_PATH, width=110)

        st.markdown(
            """
            <h2 style='text-align:center;font-size:24px;font-weight:800;'>
                Acceso al sistema
            </h2>
            <p style='text-align:center;opacity:.75;'>
                Ingrese con su cuenta de administrador o docente.
            </p>
            """,
            unsafe_allow_html=True
        )

        credenciales = cargar_credenciales_admin()

        if (
            credenciales.get("requiere_cambio_password")
            and st.session_state.password_inicial_temporal
        ):
            st.warning(
                "Primera configuración: su contraseña temporal de administrador es:"
            )
            st.code(
                st.session_state.password_inicial_temporal,
                language=None
            )
            st.caption(
                "Después de iniciar sesión, cámbiela desde ⚙️ Mi cuenta."
            )

        with st.form("form_login_sistema"):
            usuario = st.text_input(
                "Usuario",
                placeholder="Ej. admin.nazareno o usuario docente"
            )
            password = st.text_input(
                "Contraseña",
                placeholder="Ingrese su contraseña",
                type="password"
            )
            ingresar = st.form_submit_button(
                "🔐 Ingresar",
                type="primary",
                use_container_width=True
            )

        if ingresar:
            usuario_limpio = normalizar_texto(usuario)

            if verificar_admin(usuario_limpio, password):
                st.session_state.logged_in = True
                st.session_state.rol = "admin"
                st.session_state.username = usuario_limpio
                st.session_state.pagina = "app"
                st.session_state.vista_actual = "inicio"
                st.session_state.password_inicial_temporal = None
                st.rerun()

            docente = verificar_docente(usuario_limpio, password)

            if docente:
                st.session_state.logged_in = True
                st.session_state.rol = "docente"
                st.session_state.username = usuario_limpio
                st.session_state.docente_id = docente.get("id")
                st.session_state.pagina = "app"
                st.session_state.vista_actual = "inicio_docente"
                st.session_state.password_inicial_temporal = None
                st.rerun()

            st.error("Usuario o contraseña incorrectos.")

        st.markdown("</div>", unsafe_allow_html=True)


# ============================================================
# CARNET PÚBLICO
# ============================================================
def pantalla_ver_carnet_publico(codigo):
    codigo = normalizar_codigo(codigo)

    col_a, col_b, col_c = st.columns(
        [1, 1.3, 1]
    )

    with col_b:
        if LOGO_PATH:
            lc1, lc2, lc3 = st.columns(3)

            with lc2:
                st.image(
                    LOGO_PATH,
                    width=100
                )

        datos = cargar_registro(codigo)

        if not datos:
            st.error(
                "No se encontró ningún carnet con ese código."
            )
        else:
            renderizar_carnet(datos)

        if st.button(
            "Ir al inicio",
            use_container_width=True,
            key="btn_volver_publico"
        ):
            st.query_params.clear()
            st.session_state.codigo_escaneado = None
            st.session_state.pagina = "login_admin"
            st.rerun()


# ============================================================
# DASHBOARD
# ============================================================
def panel_inicio():
    total_carnets = len(listar_registros())
    total_docentes = len(cargar_docentes())

    df_asis = cargar_asistencias()
    fecha_hoy = datetime.now().strftime(
        "%Y-%m-%d"
    )

    asistencias_hoy = int(
        (df_asis["fecha"] == fecha_hoy).sum()
    ) if not df_asis.empty else 0

    col_izq, col_der = st.columns(
        [5, 7],
        gap="large"
    )

    with col_izq:
        if LOGO_PATH:
            st.image(
                LOGO_PATH,
                width=170
            )

        st.markdown(
            """
            <h1 style='font-size:36px;
                       font-weight:800;'>
                Bienvenido de<br>
                nuevo
            </h1>
            """,
            unsafe_allow_html=True
        )

        st.markdown(
            f"""
            <p style='font-size:16px;'>
                Panel de control del sistema de carnet
                y asistencia estudiantil del
                {NOMBRE_COLEGIO_LINEA2.title()}.
            </p>
            """,
            unsafe_allow_html=True
        )

        st.caption(
            f"Versión del sistema: {APP_VERSION}"
        )

    with col_der:
        st.markdown(
            "<div class='card-panel'>",
            unsafe_allow_html=True
        )

        st.markdown(
            "### 📊 Resumen general"
        )

        m1, m2, m3 = st.columns(3)

        m1.metric(
            "🪪 Carnets registrados",
            total_carnets
        )

        m2.metric(
            "✅ Asistencias hoy",
            asistencias_hoy
        )

        m3.metric(
            "👩‍🏫 Docentes",
            total_docentes
        )

        st.markdown(
            "</div>",
            unsafe_allow_html=True
        )

        st.markdown(
            "<div class='card-panel'>",
            unsafe_allow_html=True
        )

        st.markdown(
            "### 🚀 Accesos rápidos"
        )

        b1, b2, b3 = st.columns(3)

        with b1:
            if st.button(
                "📝 Registrar carnet",
                use_container_width=True
            ):
                st.session_state.vista_actual = "registro"
                st.session_state.ultimo_codigo_registrado = None
                st.rerun()

        with b2:
            if st.button(
                "📅 Ir a asistencia",
                use_container_width=True
            ):
                st.session_state.vista_actual = "asistencia"
                st.rerun()

        with b3:
            if st.button(
                "📋 Ver carnets",
                use_container_width=True
            ):
                st.session_state.vista_actual = "datos"
                st.rerun()

        st.markdown(
            "</div>",
            unsafe_allow_html=True
        )


# ============================================================
# REGISTRO DE CARNET
# ============================================================
def pantalla_registro():
    if st.session_state.ultimo_codigo_registrado:
        st.markdown(
            "<h2 style='text-align:center;'>¡Carnet registrado con éxito! ✅</h2>",
            unsafe_allow_html=True
        )
        datos = cargar_registro(st.session_state.ultimo_codigo_registrado)
        if datos:
            c1, c2, c3 = st.columns(3)
            with c2:
                renderizar_carnet(datos)
            inscripciones = obtener_inscripciones_estudiante(datos)
            if inscripciones:
                st.markdown("### 📚 Clases del estudiante")
                st.dataframe(
                    pd.DataFrame(inscripciones)[["nivel", "grado", "materia", "docente"]],
                    use_container_width=True, hide_index=True
                )
        if st.button("Registrar otro carnet", use_container_width=True, key="btn_otro_carnet"):
            st.session_state.ultimo_codigo_registrado = None
            st.rerun()
        return

    col_izq, col_der = st.columns([5, 7], gap="large")
    with col_izq:
        if LOGO_PATH:
            st.image(LOGO_PATH, width=170)
        st.markdown(
            """
            <h1 style='font-size:38px;font-weight:800;'>Registro de<br>carnet estudiantil</h1>
            <p style='font-size:16px;'>Registre al estudiante y asígnelo a una o varias clases.
            El mismo QR podrá utilizarse para registrar asistencia en distintas materias.</p>
            """, unsafe_allow_html=True
        )

    with col_der:
        st.markdown("<div class='card-panel'>", unsafe_allow_html=True)
        st.markdown("### 👤 Datos personales")
        c1, c2 = st.columns(2)
        with c1:
            nombres = st.text_input("Nombres *", placeholder="Ingrese nombres", key="registro_nombres")
        with c2:
            apellidos = st.text_input("Apellidos *", placeholder="Ingrese apellidos", key="registro_apellidos")
        c3, c4 = st.columns(2)
        with c3:
            dpi = st.text_input("DPI", placeholder="Opcional", key="registro_dpi")
        with c4:
            fecha_nac = st.date_input("Fecha de nacimiento *", value=None,
                                      min_value=date(1995, 1, 1), max_value=date.today(), key="registro_fecha_nac")

        st.markdown("### 📝 Información académica")
        c5, c6 = st.columns(2)
        with c5:
            carnet_actual = st.text_input("No. de carnet actual (si tiene)", placeholder="Opcional", key="registro_carnet_actual")
        with c6:
            nivel_seleccionado = st.selectbox("Nivel educativo *", ["Seleccione el nivel"] + NIVELES_EDUCATIVOS,
                                              key="registro_nivel")

        st.markdown("### 📚 Inscripción a clases")
        todas_clases = obtener_clases_docentes()
        clases_nivel_disponibles = (
            [c for c in todas_clases if c.get("nivel") == nivel_seleccionado]
            if nivel_seleccionado != "Seleccione el nivel" else []
        )

        # Mostrar automáticamente los docentes registrados en el nivel elegido.
        docentes_nivel = sorted({
            normalizar_texto(c.get("docente", ""))
            for c in clases_nivel_disponibles if c.get("docente")
        })
        if nivel_seleccionado == "Seleccione el nivel":
            st.info("Seleccione primero el nivel educativo para mostrar sus docentes y clases.")
        elif docentes_nivel:
            st.selectbox(
                "👩‍🏫 Docente(s) registrado(s) en este nivel",
                docentes_nivel,
                index=0,
                disabled=True,
                key="docente_segun_nivel"
            )
            if len(docentes_nivel) > 1:
                st.caption("Docentes disponibles en el nivel: " + ", ".join(docentes_nivel))
        else:
            st.warning("No hay docentes registrados en este nivel educativo.")

        opciones = {"Sin asignación por ahora": None}
        for clase in clases_nivel_disponibles:
            opciones[etiqueta_clase(clase)] = clase

        seleccionadas = st.multiselect(
            "Seleccione una o varias clases del estudiante *",
            options=list(opciones.keys()),
            key="registro_clases",
            help="Solo se muestran clases y docentes del nivel educativo seleccionado."
        )
        clases_seleccionadas = [opciones[x] for x in seleccionadas if opciones[x] is not None]

        if clases_seleccionadas:
            st.dataframe(
                pd.DataFrame(clases_seleccionadas)[["nivel", "grado", "materia", "docente"]],
                use_container_width=True, hide_index=True
            )

        st.markdown("### 📷 Fotografía")
        foto_estudiante = st.file_uploader("Foto del estudiante *", type=["png", "jpg", "jpeg"], key="registro_foto")
        if foto_estudiante is not None:
            st.image(foto_estudiante, width=140)

        if st.button("💾 Registrar carnet", type="primary", use_container_width=True, key="btn_registrar_carnet"):
            nombres_limpios = normalizar_texto(nombres)
            apellidos_limpios = normalizar_texto(apellidos)
            if not nombres_limpios or not apellidos_limpios:
                st.error("Ingrese Nombres y Apellidos."); return
            if nivel_seleccionado == "Seleccione el nivel":
                st.error("Seleccione el nivel educativo."); return
            if fecha_nac is None:
                st.error("Ingrese la fecha de nacimiento."); return
            if foto_estudiante is None:
                st.error("Suba una foto del estudiante."); return
            if not clases_seleccionadas and todas_clases:
                st.error("Seleccione al menos una clase para el estudiante."); return
            if any(c.get("nivel") != nivel_seleccionado for c in clases_seleccionadas):
                st.error("Todas las clases seleccionadas deben pertenecer al nivel educativo elegido."); return

            dpi_final = normalizar_texto(dpi) if normalizar_texto(dpi) else "No aplica"
            duplicado = buscar_duplicado_por_dpi(dpi_final) if dpi_final != "No aplica" else None
            if duplicado:
                st.error(f"Este DPI ya está registrado con el código {duplicado.get('codigo', '')}."); return

            codigo = normalizar_codigo(carnet_actual) if normalizar_texto(carnet_actual) else generar_codigo_aleatorio()
            if not codigo_valido(codigo):
                st.error("El número de carnet debe tener entre 3 y 30 caracteres y usar solo letras, números, guion o guion bajo."); return
            if codigo_existente(codigo):
                st.error(f"Ya existe un carnet con el código {codigo}. Use otro código."); return

            ext_foto = os.path.splitext(foto_estudiante.name)[1].lower().replace(".", "")
            if ext_foto not in {"jpg", "jpeg", "png"}:
                st.error("Formato de foto no permitido."); return
            contenido_foto = foto_estudiante.getvalue()
            if len(contenido_foto) > 5 * 1024 * 1024:
                st.error("La foto supera el límite de 5 MB."); return

            nombre_foto = nombre_archivo_seguro(codigo, "foto", ext_foto)
            ruta_foto = os.path.join(OS_DIR, nombre_foto)
            primera = clases_seleccionadas[0] if clases_seleccionadas else {
                "nivel": nivel_seleccionado, "grado": "Sin grado", "materia": "No asignada", "docente": "No asignado"
            }
            clases_guardadas = []
            for c in clases_seleccionadas:
                clases_guardadas.append(dict(c))
            datos = {
                "nombres": nombres_limpios, "apellidos": apellidos_limpios, "dpi": dpi_final,
                "fecha_nacimiento": str(fecha_nac), "codigo": codigo,
                "grado": primera.get("grado", "Sin grado"), "nivel": nivel_seleccionado,
                "docente": primera.get("docente", "No asignado"), "materia": primera.get("materia", "No asignada"),
                "clases": clases_guardadas,
            }
            try:
                with open(ruta_foto, "wb") as f: f.write(contenido_foto)
                datos["foto"] = nombre_foto
                ruta_qr = generar_qr_carnet(codigo)
                escribir_json_atomico(ruta_registro(codigo), datos)
            except OSError as exc:
                eliminar_archivo_seguro(ruta_foto)
                if 'ruta_qr' in locals(): eliminar_archivo_seguro(ruta_qr)
                st.error(f"No se pudo guardar el registro: {exc}"); return

            st.session_state.ultimo_codigo_registrado = codigo
            st.rerun()
        st.markdown("</div>", unsafe_allow_html=True)


# ============================================================
# DOCENTES# ============================================================
# DOCENTES
# ============================================================
def construir_clases_desde_filas(filas):
    clases = []
    vistas = set()
    for nivel, grado, materia in filas:
        nivel = normalizar_texto(nivel)
        grado = normalizar_texto(grado)
        materia = normalizar_texto(materia)
        if nivel and grado and materia:
            clave = (nivel.casefold(), grado.casefold(), materia.casefold())
            if clave not in vistas:
                clases.append({"nivel": nivel, "grado": grado, "materia": materia})
                vistas.add(clave)
    return clases


def render_filas_clases_docente(prefix, clases_iniciales=None):
    clases_iniciales = clases_iniciales or []
    filas = []
    for i in range(MAX_CLASES_DOCENTE):
        actual = clases_iniciales[i] if i < len(clases_iniciales) else {}
        nivel_actual = actual.get("nivel", "")
        nivel_op = ["— Sin clase —"] + NIVELES_EDUCATIVOS
        if nivel_actual not in NIVELES_EDUCATIVOS:
            nivel_actual = "— Sin clase —"
        c1, c2, c3 = st.columns([2.1, 3.2, 3.2])
        with c1:
            nivel = st.selectbox("Nivel" if i == 0 else "", nivel_op,
                                 index=nivel_op.index(nivel_actual), key=f"{prefix}_nivel_{i}")
        grados = GRADOS_SECCIONES_POR_NIVEL.get(nivel, [])
        grado_actual = actual.get("grado", "")
        grado_op = ["— Seleccione grado/sección —"] + grados
        if grado_actual and grado_actual not in grados:
            grado_op.append(grado_actual)
        if nivel == "— Sin clase —":
            grado_actual_sel = "— Seleccione grado/sección —"
        else:
            grado_actual_sel = grado_actual if grado_actual in grado_op else "— Seleccione grado/sección —"
        with c2:
            grado = st.selectbox("Grado / Sección" if i == 0 else "", grado_op,
                                 index=grado_op.index(grado_actual_sel), key=f"{prefix}_grado_{i}")
        with c3:
            materia = st.text_input("Materia" if i == 0 else "", value=actual.get("materia", ""),
                                    placeholder="Ej. Matemáticas", key=f"{prefix}_materia_{i}")
        if nivel != "— Sin clase —" and grado != "— Seleccione grado/sección —" and normalizar_texto(materia):
            filas.append((nivel, grado, materia))
    return construir_clases_desde_filas(filas)


def panel_docentes():
    st.markdown("### 👩‍🏫 Gestión de docentes")
    st.caption("Cree docentes y asígneles directamente las materias y grados/secciones que pueden atender.")
    docentes = cargar_docentes()

    if st.session_state.docente_editando is not None:
        doc_edit = next((d for d in docentes if d.get("id") == st.session_state.docente_editando), None)
        if doc_edit:
            st.markdown("#### ✏️ Editar perfil docente")
            c1, c2 = st.columns(2)
            with c1:
                nombre_edit = st.text_input("Nombre completo *", value=doc_edit.get("nombre", ""), key="ed_doc_nombre")
                usuario_edit = st.text_input("Usuario de acceso *", value=doc_edit.get("usuario", ""), key="ed_doc_usuario")
                nueva_password = st.text_input("Nueva contraseña (opcional)", type="password", placeholder="Mínimo 8 caracteres", key="ed_doc_pass")
                confirmar_password = st.text_input("Confirmar nueva contraseña", type="password", key="ed_doc_pass2")
            with c2:
                niveles_edit = st.multiselect("Nivel(es) educativo(s) *", NIVELES_EDUCATIVOS, default=doc_edit.get("niveles", []), key="ed_doc_niveles")
            st.markdown("**Clases asignadas**")
            clases_edit = render_filas_clases_docente("edit_doc", doc_edit.get("clases", []))
            ce1, ce2 = st.columns(2)
            with ce1:
                guardar_edit = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="guardar_doc_edit")
            with ce2:
                cancelar_edit = st.button("Cancelar", use_container_width=True, key="cancelar_doc_edit")
            if cancelar_edit:
                st.session_state.docente_editando = None; st.rerun()
            if guardar_edit:
                nombre_limpio=normalizar_texto(nombre_edit); usuario_limpio=normalizar_texto(usuario_edit).lower()
                usuario_repetido=any(d.get("id") != doc_edit.get("id") and normalizar_texto(d.get("usuario", "")).casefold()==usuario_limpio.casefold() for d in docentes)
                if not nombre_limpio: st.error("El nombre no puede estar vacío.")
                elif not re.fullmatch(r"[A-Za-z0-9._-]{3,40}", usuario_limpio): st.error("El usuario debe tener entre 3 y 40 caracteres y usar letras, números, punto, guion o guion bajo.")
                elif usuario_repetido: st.error("Ese usuario ya pertenece a otro docente.")
                elif not niveles_edit: st.error("Seleccione al menos un nivel.")
                elif not clases_edit: st.error("Registre al menos una clase.")
                elif any(c.get("nivel") not in niveles_edit for c in clases_edit): st.error("Todas las clases deben pertenecer a los niveles seleccionados.")
                elif nueva_password and len(nueva_password)<8: st.error("La nueva contraseña debe tener al menos 8 caracteres.")
                elif nueva_password and nueva_password!=confirmar_password: st.error("Las contraseñas no coinciden.")
                else:
                    for c in clases_edit:
                        c.update({"id":crear_clase_id(doc_edit.get("id"),c["grado"],c["materia"]),"docente_id":doc_edit.get("id"),"docente":nombre_limpio})
                    doc_edit.update({"nombre":nombre_limpio,"usuario":usuario_limpio,"niveles":niveles_edit,"clases":clases_edit,
                                     "materias":sorted({c["materia"] for c in clases_edit},key=str.casefold),"grados":sorted({c["grado"] for c in clases_edit},key=str.casefold)})
                    if nueva_password: doc_edit["password_hash"]=hash_password(nueva_password); doc_edit["requiere_cambio_password"]=False
                    guardar_docentes(docentes); st.session_state.docente_editando=None; st.success("Perfil docente actualizado correctamente."); st.rerun()
            return

    with st.expander("➕ Crear cuenta de nuevo docente", expanded=True):
        c1,c2=st.columns(2)
        with c1:
            nombre_docente=st.text_input("Nombre completo *",placeholder="Ej. Prof. Juan Carlos López",key="nuevo_doc_nombre")
            usuario_docente=st.text_input("Usuario *",placeholder="Ej. juan.lopez",key="nuevo_doc_usuario")
            password_docente=st.text_input("Contraseña inicial *",type="password",placeholder="Mínimo 8 caracteres",key="nuevo_doc_pass")
            confirmar_docente=st.text_input("Confirmar contraseña *",type="password",key="nuevo_doc_pass2")
        with c2:
            niveles_docente=st.multiselect("Nivel(es) educativo(s) *",NIVELES_EDUCATIVOS,key="nuevo_doc_niveles")
        st.markdown("**Clases asignadas**")
        clases=render_filas_clases_docente("nuevo_doc")
        st.caption("Seleccione el nivel, luego el grado/sección del catálogo y escriba la materia. Puede registrar hasta 12 clases por docente.")
        guardar=st.button("🔐 Crear cuenta docente",type="primary",use_container_width=True,key="guardar_nuevo_doc")
        if guardar:
            nombre_limpio=normalizar_texto(nombre_docente); usuario_limpio=normalizar_texto(usuario_docente).lower()
            usuario_admin=normalizar_texto(cargar_credenciales_admin().get("usuario","")).casefold()
            usuario_repetido=usuario_limpio.casefold()==usuario_admin or any(normalizar_texto(d.get("usuario","")).casefold()==usuario_limpio.casefold() for d in docentes)
            nombre_repetido=any(normalizar_texto(d.get("nombre","")).casefold()==nombre_limpio.casefold() for d in docentes)
            if not nombre_limpio: st.error("Ingrese el nombre del docente.")
            elif not re.fullmatch(r"[A-Za-z0-9._-]{3,40}",usuario_limpio): st.error("El usuario debe tener entre 3 y 40 caracteres y usar letras, números, punto, guion o guion bajo.")
            elif usuario_repetido: st.error("Ese usuario ya está en uso.")
            elif nombre_repetido: st.error("Ya existe un docente con ese nombre.")
            elif not niveles_docente: st.error("Seleccione al menos un nivel.")
            elif len(password_docente)<8: st.error("La contraseña debe tener al menos 8 caracteres.")
            elif password_docente!=confirmar_docente: st.error("Las contraseñas no coinciden.")
            elif not clases: st.error("Registre al menos una clase.")
            elif any(c.get("nivel") not in niveles_docente for c in clases): st.error("Todas las clases deben pertenecer a los niveles seleccionados.")
            else:
                docente_id=uuid.uuid4().hex[:12]
                for c in clases: c.update({"id":crear_clase_id(docente_id,c["grado"],c["materia"]),"docente_id":docente_id,"docente":nombre_limpio})
                nuevo={"id":docente_id,"nombre":nombre_limpio,"usuario":usuario_limpio,"password_hash":hash_password(password_docente),"requiere_cambio_password":False,
                       "materias":sorted({c["materia"] for c in clases},key=str.casefold),"niveles":niveles_docente,"grados":sorted({c["grado"] for c in clases},key=str.casefold),"clases":clases}
                docentes.append(nuevo); guardar_docentes(docentes); st.success(f"Cuenta de {nombre_limpio} creada correctamente."); st.rerun()

    st.markdown("---"); st.markdown("### 📋 Docentes registrados")
    if not docentes: st.info("No hay docentes registrados."); return
    for idx,doc in enumerate(docentes):
        with st.container():
            cd1,cd2,cd3=st.columns([3,4.5,2])
            with cd1:
                st.markdown(f"**👨‍🏫 {doc.get('nombre','')}**"); st.caption(f"Usuario: {doc.get('usuario','Sin cuenta')}"); st.caption("Niveles: "+", ".join(doc.get("niveles",[])))
            with cd2:
                clases_doc=doc.get("clases",[]); st.caption("📚 "+" · ".join(f"{c.get('grado')} — {c.get('materia')}" for c in clases_doc) if clases_doc else "Sin clases asignadas.")
            with cd3:
                cbtn1,cbtn2=st.columns(2)
                with cbtn1:
                    if st.button("✏️ Editar",key=f"edit_doc_{doc.get('id')}",use_container_width=True): st.session_state.docente_editando=doc.get("id"); st.rerun()
                with cbtn2:
                    if st.button("🗑️ Eliminar",key=f"del_doc_{doc.get('id')}",use_container_width=True): docentes.pop(idx); guardar_docentes(docentes); st.rerun()
            st.divider()


# ============================================================
# CARNETS REGISTRADOS# ============================================================
# CARNETS REGISTRADOS
# ============================================================
def eliminar_archivo_seguro(path):
    if path and os.path.exists(path):
        try:
            os.remove(path)
        except OSError:
            pass


def generar_qr_carnet(codigo):
    codigo = normalizar_codigo(codigo)

    url_destinatario = (
        f"{BASE_URL}/"
        f"?codigo="
        f"{urlparse.quote(codigo)}"
    )

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3
    )

    qr.add_data(url_destinatario)
    qr.make(fit=True)

    img_qr = qr.make_image(
        fill_color="black",
        back_color="white"
    )

    ruta_qr = os.path.join(
        OS_DIR,
        f"{codigo}_qr.png"
    )

    img_qr.save(ruta_qr)
    return ruta_qr


def actualizar_asistencias_por_cambio_carnet(codigo_anterior, datos_actualizados):
    df = cargar_asistencias()
    if df.empty: return
    codigo_anterior = normalizar_codigo(codigo_anterior); codigo_nuevo = normalizar_codigo(datos_actualizados.get("codigo", ""))
    mascara = df["codigo"].astype(str).str.upper() == codigo_anterior
    if not mascara.any(): return
    df.loc[mascara, "codigo"] = codigo_nuevo
    df.loc[mascara, "nombres"] = datos_actualizados.get("nombres", "")
    df.loc[mascara, "apellidos"] = datos_actualizados.get("apellidos", "")
    # La clase histórica se conserva; editar la inscripción actual no cambia el pasado.
    guardar_asistencias(df)


def guardar_cambios_carnet(
    codigo_anterior,
    datos_originales,
    datos_actualizados,
    nueva_foto=None
):
    codigo_anterior = normalizar_codigo(
        codigo_anterior
    )
    codigo_nuevo = normalizar_codigo(
        datos_actualizados.get("codigo", "")
    )

    if not codigo_valido(codigo_nuevo):
        raise ValueError(
            "El número de carnet debe tener entre 3 y 30 "
            "caracteres y usar solo letras, números, guion "
            "o guion bajo."
        )

    ruta_anterior = ruta_registro(
        codigo_anterior
    )
    ruta_nueva = ruta_registro(
        codigo_nuevo
    )

    if (
        codigo_nuevo != codigo_anterior
        and os.path.exists(ruta_nueva)
    ):
        raise ValueError(
            f"Ya existe un carnet con el código "
            f"{codigo_nuevo}."
        )

    foto_anterior = datos_originales.get("foto")
    ruta_foto_anterior = (
        os.path.join(
            OS_DIR,
            os.path.basename(foto_anterior)
        )
        if foto_anterior
        else None
    )

    foto_nueva_guardada = None
    ruta_foto_nueva = None

    try:
        if nueva_foto is not None:
            ext_foto = (
                os.path.splitext(
                    nueva_foto.name
                )[1]
                .lower()
                .replace(".", "")
            )

            if ext_foto not in {
                "jpg",
                "jpeg",
                "png"
            }:
                raise ValueError(
                    "Formato de foto no permitido."
                )

            contenido_foto = (
                nueva_foto.getvalue()
            )

            if len(contenido_foto) > 5 * 1024 * 1024:
                raise ValueError(
                    "La foto supera el límite de 5 MB."
                )

            foto_nueva_guardada = (
                nombre_archivo_seguro(
                    codigo_nuevo,
                    "foto",
                    ext_foto
                )
            )
            ruta_foto_nueva = os.path.join(
                OS_DIR,
                foto_nueva_guardada
            )

            with open(
                ruta_foto_nueva,
                "wb"
            ) as f:
                f.write(contenido_foto)

            datos_actualizados["foto"] = (
                foto_nueva_guardada
            )

        elif foto_anterior:
            extension_anterior = (
                os.path.splitext(
                    foto_anterior
                )[1]
                .lower()
                .replace(".", "")
            )

            if (
                codigo_nuevo != codigo_anterior
                and os.path.exists(
                    ruta_foto_anterior
                )
            ):
                foto_nueva_guardada = (
                    nombre_archivo_seguro(
                        codigo_nuevo,
                        "foto",
                        extension_anterior
                    )
                )
                ruta_foto_nueva = os.path.join(
                    OS_DIR,
                    foto_nueva_guardada
                )

                os.replace(
                    ruta_foto_anterior,
                    ruta_foto_nueva
                )

                datos_actualizados["foto"] = (
                    foto_nueva_guardada
                )

        ruta_qr_anterior = os.path.join(
            OS_DIR,
            f"{codigo_anterior}_qr.png"
        )
        ruta_qr_nueva = generar_qr_carnet(
            codigo_nuevo
        )

        escribir_json_atomico(
            ruta_nueva,
            datos_actualizados
        )

        if (
            codigo_nuevo != codigo_anterior
            and os.path.exists(ruta_anterior)
        ):
            eliminar_archivo_seguro(
                ruta_anterior
            )

        if (
            ruta_qr_anterior != ruta_qr_nueva
            and os.path.exists(ruta_qr_anterior)
        ):
            eliminar_archivo_seguro(
                ruta_qr_anterior
            )

        if (
            nueva_foto is not None
            and foto_anterior
            and ruta_foto_anterior
            and ruta_foto_anterior != ruta_foto_nueva
        ):
            eliminar_archivo_seguro(
                ruta_foto_anterior
            )

        actualizar_asistencias_por_cambio_carnet(
            codigo_anterior,
            datos_actualizados
        )

    except Exception:
        if ruta_foto_nueva and (
            nueva_foto is not None
        ):
            eliminar_archivo_seguro(
                ruta_foto_nueva
            )

        if (
            codigo_nuevo != codigo_anterior
            and os.path.exists(ruta_nueva)
        ):
            eliminar_archivo_seguro(
                ruta_nueva
            )

        raise


def panel_editar_carnet():
    codigo_editar = st.session_state.carnet_editando
    if not codigo_editar: return False
    datos = cargar_registro(codigo_editar)
    if not datos:
        st.error("No se encontró el carnet seleccionado."); st.session_state.carnet_editando = None; return False

    st.markdown("### ✏️ Editar carnet estudiantil")
    st.info("Puede cambiar las clases del estudiante. El mismo QR seguirá siendo válido para todas sus clases.")
    docentes_registrados = cargar_docentes()
    todas_clases = obtener_clases_docentes(docentes_registrados)
    opciones = {etiqueta_clase(c): c for c in todas_clases}
    actuales = obtener_inscripciones_estudiante(datos)
    seleccion_inicial = []
    for c in actuales:
        etiqueta = etiqueta_clase(c)
        if etiqueta not in opciones: opciones[etiqueta] = c
        seleccion_inicial.append(etiqueta)

    fecha_original = datos.get("fecha_nacimiento")
    try: fecha_original = datetime.strptime(str(fecha_original), "%Y-%m-%d").date()
    except (ValueError, TypeError): fecha_original = date(max(1995, date.today().year - 18), 1, 1)

    with st.form("form_editar_carnet_v4"):
        st.markdown("#### 👤 Datos personales")
        c1, c2 = st.columns(2)
        with c1: nombres = st.text_input("Nombres *", value=str(datos.get("nombres", "")))
        with c2: apellidos = st.text_input("Apellidos *", value=str(datos.get("apellidos", "")))
        c3, c4 = st.columns(2)
        with c3: dpi = st.text_input("DPI", value="" if str(datos.get("dpi", "")).casefold() == "no aplica" else str(datos.get("dpi", "")))
        with c4: fecha_nac = st.date_input("Fecha de nacimiento *", value=fecha_original, min_value=date(1995,1,1), max_value=date.today())
        st.markdown("#### 📝 Información académica")
        niveles_opciones = ["Seleccione el nivel"] + NIVELES_EDUCATIVOS
        nivel_actual = datos.get("nivel", "Seleccione el nivel")
        if nivel_actual not in niveles_opciones: nivel_actual = "Seleccione el nivel"
        c5, c6 = st.columns(2)
        with c5: codigo = st.text_input("No. de carnet *", value=str(datos.get("codigo", "")))
        with c6: nivel_seleccionado = st.selectbox("Nivel educativo *", niveles_opciones, index=niveles_opciones.index(nivel_actual))
        seleccionadas = st.multiselect("Clases / materias del estudiante *", list(opciones.keys()), default=[x for x in seleccion_inicial if x in opciones], help="Puede seleccionar varias materias de distintos docentes.")
        st.markdown("#### 📷 Foto")
        foto_actual = datos.get("foto")
        if foto_actual:
            ruta_foto_actual = os.path.join(OS_DIR, os.path.basename(foto_actual))
            if os.path.exists(ruta_foto_actual): st.image(ruta_foto_actual, width=140)
        nueva_foto = st.file_uploader("Nueva foto (opcional)", type=["png","jpg","jpeg"])
        col_guardar, col_cancelar = st.columns(2)
        with col_guardar: guardar = st.form_submit_button("💾 Guardar cambios", type="primary", use_container_width=True)
        with col_cancelar: cancelar = st.form_submit_button("Cancelar", use_container_width=True)

    if cancelar:
        st.session_state.carnet_editando = None; st.rerun()
    if guardar:
        nombres_limpios = normalizar_texto(nombres); apellidos_limpios = normalizar_texto(apellidos); codigo_limpio = normalizar_codigo(codigo)
        clases_seleccionadas = [opciones[x] for x in seleccionadas if opciones[x] is not None]
        if not nombres_limpios or not apellidos_limpios: st.error("Ingrese Nombres y Apellidos."); return True
        if nivel_seleccionado == "Seleccione el nivel": st.error("Seleccione el nivel educativo."); return True
        if not clases_seleccionadas and todas_clases: st.error("Seleccione al menos una clase."); return True
        if any(c.get("nivel") != nivel_seleccionado for c in clases_seleccionadas): st.error("Todas las clases deben pertenecer al nivel educativo elegido."); return True
        if not codigo_valido(codigo_limpio): st.error("El número de carnet debe tener entre 3 y 30 caracteres y usar solo letras, números, guion o guion bajo."); return True
        registros = listar_registros()
        if any(normalizar_codigo(r.get("codigo", "")) == codigo_limpio and normalizar_codigo(r.get("codigo", "")) != normalizar_codigo(codigo_editar) for r in registros): st.error(f"Ya existe otro carnet con el código {codigo_limpio}."); return True
        dpi_final = normalizar_texto(dpi) if normalizar_texto(dpi) else "No aplica"
        if dpi_final.casefold() != "no aplica":
            dn = dpi_final.lower().replace(" ", "")
            for r in registros:
                if normalizar_codigo(r.get("codigo", "")) == normalizar_codigo(codigo_editar): continue
                if normalizar_texto(r.get("dpi", "")).lower().replace(" ", "") == dn and dn: st.error(f"Este DPI ya está registrado con el código {r.get('codigo', '')}."); return True
        primera = clases_seleccionadas[0] if clases_seleccionadas else {"grado":"Sin grado","docente":"No asignado","materia":"No asignada"}
        datos_actualizados = dict(datos)
        datos_actualizados.update({"nombres":nombres_limpios,"apellidos":apellidos_limpios,"dpi":dpi_final,"fecha_nacimiento":str(fecha_nac),"codigo":codigo_limpio,
                                   "grado":primera.get("grado","Sin grado"),"nivel":nivel_seleccionado,"docente":primera.get("docente","No asignado"),"materia":primera.get("materia","No asignada"),"clases":[dict(c) for c in clases_seleccionadas]})
        try: guardar_cambios_carnet(codigo_editar, datos, datos_actualizados, nueva_foto)
        except (ValueError, OSError, json.JSONDecodeError) as exc: st.error(f"No se pudieron guardar los cambios. {exc}"); return True
        st.session_state.carnet_editando = None; st.session_state.ultimo_codigo_registrado = codigo_limpio
        st.success("Carnet actualizado exitosamente. El QR también fue actualizado."); st.rerun()
    return True


def panel_datos():
    if st.session_state.carnet_editando:
        if panel_editar_carnet():
            return

    st.markdown(
        "### 📋 Carnets registrados"
    )

    registros = listar_registros()

    if not registros:
        st.info(
            "Aún no hay carnets registrados."
        )
        return

    busqueda = st.text_input(
        "Buscar por nombre, apellido, DPI o código",
        placeholder="Escriba para filtrar...",
        key="buscar_carnet"
    )

    if busqueda:
        termino = normalizar_texto(
            busqueda
        ).casefold()

        registros = [
            r for r in registros
            if termino in (
                f"{r.get('codigo', '')} "
                f"{r.get('nombres', '')} "
                f"{r.get('apellidos', '')} "
                f"{r.get('dpi', '')} "
                f"{r.get('grado', '')}"
            ).casefold()
        ]

    if not registros:
        st.warning(
            "No se encontraron estudiantes."
        )
        return

    grados_presentes = {
        normalizar_grado(
            r.get("grado")
        )
        for r in registros
    }

    for grado_actual in ordenar_grados_presentes(
        grados_presentes
    ):
        registros_grado = [
            r for r in registros
            if normalizar_grado(
                r.get("grado")
            ) == grado_actual
        ]

        with st.expander(
            f"🎓 {grado_actual} — "
            f"{len(registros_grado)} estudiante(s)",
            expanded=True
        ):
            df_grupo = pd.DataFrame(
                registros_grado
            )

            mostrar_tabla_con_carnet(
                df_grupo,
                key_tabla=(
                    "tabla_grado_"
                    + re.sub(
                        r"[^A-Za-z0-9_-]",
                        "_",
                        grado_actual
                    )
                ),
                permitir_edicion=True
            )


# ============================================================
# ASISTENCIA
# ============================================================
def extraer_codigo_desde_qr(valor_qr):
    valor_qr = normalizar_texto(valor_qr)
    if not valor_qr: return None
    try:
        parsed = urlparse.urlparse(valor_qr)
        query_codigo = urlparse.parse_qs(parsed.query).get("codigo", [None])[0]
        if query_codigo: return normalizar_codigo(query_codigo)
    except Exception: pass
    return normalizar_codigo(valor_qr)


def clases_permitidas_para_usuario():
    if st.session_state.get("rol") == "admin": return obtener_clases_docentes()
    docente = docente_actual()
    if not docente: return []
    clases = []
    for c in docente.get("clases", []):
        cc = dict(c); cc["docente_id"] = docente.get("id"); cc["docente"] = docente.get("nombre", "")
        cc["nivel"] = cc.get("nivel") or (docente.get("niveles", [""])[0] if docente.get("niveles") else "")
        cc["id"] = cc.get("id") or crear_clase_id(docente.get("id"), cc.get("grado", ""), cc.get("materia", ""))
        clases.append(cc)
    return clases


def procesar_qr_asistencia(clase):
    st.markdown("#### 📸 Escanear QR")
    foto_cam = st.camera_input("Apunte el código QR del carnet a la cámara", key=f"camara_qr_{clase.get('id')}")
    if foto_cam is not None:
        cv2_img = cv2.imdecode(np.frombuffer(foto_cam.getvalue(), np.uint8), cv2.IMREAD_COLOR)
        if cv2_img is None:
            st.error("No se pudo leer la imagen."); return
        detector = cv2.QRCodeDetector()
        try: val, pts, qr_code = detector.detectAndDecode(cv2_img)
        except cv2.error: val = ""
        codigo = extraer_codigo_desde_qr(val) if val else None
        if not codigo or not codigo_valido(codigo):
            st.error("No se pudo detectar un código QR válido."); return
        exito, msg = registrar_asistencia_codigo(codigo, clase)
        if exito: st.success(msg); st.rerun()
        else: st.warning(msg)


def panel_asistencia():
    rol = st.session_state.get("rol")
    titulo = "### 📅 Asistencia por clase" if rol == "admin" else "### 📅 Registrar asistencia"
    st.markdown(titulo)
    st.caption("Seleccione la clase primero. El mismo carnet QR puede registrar asistencia en distintas materias el mismo día.")
    clases = clases_permitidas_para_usuario()
    if not clases:
        st.info("No hay clases asignadas disponibles."); return
    opciones = {etiqueta_clase(c): c for c in clases}
    etiqueta = st.selectbox("📚 Clase / materia", list(opciones.keys()), key="asistencia_clase_seleccionada")
    clase = opciones[etiqueta]
    st.success(f"Clase activa: **{clase.get('materia')}** · {clase.get('grado')} · {clase.get('docente')}")

    col_cam, col_manual = st.columns(2)
    with col_cam:
        procesar_qr_asistencia(clase)
    with col_manual:
        st.markdown("#### 🔍 Registro manual")
        registros = [r for r in listar_registros() if estudiante_en_clase(r, clase)]
        if not registros:
            st.info("No hay estudiantes inscritos en esta clase.")
        else:
            opciones_est = {f"{r.get('codigo')} - {r.get('nombres')} {r.get('apellidos')}": r.get('codigo') for r in registros}
            est = st.selectbox("Seleccione estudiante", ["-- Seleccione --"] + sorted(opciones_est), key="asistencia_manual_estudiante")
            if st.button("Registrar asistencia", type="primary", use_container_width=True, key="btn_asistencia_manual_clase"):
                codigo = opciones_est.get(est)
                if codigo:
                    exito, msg = registrar_asistencia_codigo(codigo, clase)
                    if exito: st.success(msg); st.rerun()
                    else: st.warning(msg)
                else: st.error("Seleccione un estudiante válido.")

    st.markdown("---")
    df = cargar_asistencias()
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    if not df.empty:
        df_hoy = df[(df["fecha"] == fecha_hoy) & (df["clase_id"] == str(clase.get("id")))].reset_index(drop=True)
    else: df_hoy = pd.DataFrame(columns=COLUMNAS_ASISTENCIA)
    st.markdown(f"#### ✅ Asistencias de hoy — {len(df_hoy)}")
    if df_hoy.empty: st.info("Todavía no hay asistencias registradas para esta clase hoy.")
    else:
        mostrar_tabla_con_carnet(df_hoy, key_tabla=f"tabla_asistencia_clase_{clase.get('id')}")
        excel_bytes = generar_excel_bytes(
            df_hoy,
            titulo_reporte=f"ASISTENCIA {clase.get('materia')} - {clase.get('grado')}"
        )
        if excel_bytes is not None:
            st.download_button(
                "📊 Descargar Excel profesional",
                data=excel_bytes,
                file_name=f"asistencia_{fecha_hoy}_{clase.get('id')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary",
                key=f"btn_excel_hoy_{clase.get('id')}"
            )


# ============================================================
# HISTORIAL# ============================================================
# HISTORIAL
# ============================================================
def panel_historial_asistencia():
    st.markdown(
        "### 🗓️ Historial de asistencia"
    )

    anio_actual = datetime.now().year
    df = cargar_asistencias()

    # Los docentes solo pueden descargar y consultar la asistencia de sus clases.
    if st.session_state.get("rol") == "docente":
        ids_permitidos = {
            str(c.get("id")) for c in clases_permitidas_para_usuario()
        }
        if not df.empty:
            df = df[df["clase_id"].astype(str).isin(ids_permitidos)].copy()

    if df.empty:
        st.info(
            "Aún no hay asistencias registradas."
        )
        return

    df = df.copy()

    df["fecha_dt"] = pd.to_datetime(
        df["fecha"],
        errors="coerce"
    )

    df_anio = (
        df[
            df["fecha_dt"].dt.year
            == anio_actual
        ]
        .sort_values(
            ["fecha_dt", "hora"]
        )
    )

    if df_anio.empty:
        st.info(
            f"No hay asistencias registradas "
            f"en {anio_actual}."
        )
        return

    df_anio = df_anio.assign(
        mes_num=df_anio["fecha_dt"].dt.month
    )

    meses_presentes = sorted(
        df_anio["mes_num"].dropna().unique(),
        reverse=True
    )

    for mes_num in meses_presentes:
        mes_num = int(mes_num)

        df_mes = (
            df_anio[
                df_anio["mes_num"]
                == mes_num
            ]
            .drop(
                columns=["fecha_dt", "mes_num"]
            )
            .reset_index(drop=True)
        )

        nombre_mes = MESES_ES.get(
            mes_num,
            str(mes_num)
        )

        with st.expander(
            f"📅 {nombre_mes} {anio_actual} — "
            f"{len(df_mes)} registro(s)",
            expanded=True
        ):
            grados_presentes_mes = {
                normalizar_grado(g)
                for g in df_mes["grado"].tolist()
            }

            for grado_actual in ordenar_grados_presentes(
                grados_presentes_mes
            ):
                df_grupo_mes = df_mes[
                    df_mes["grado"].apply(
                        normalizar_grado
                    )
                    == grado_actual
                ]

                st.markdown(
                    f"**🎓 {grado_actual}**"
                )

                mostrar_tabla_con_carnet(
                    df_grupo_mes,
                    key_tabla=(
                        "tabla_historial_"
                        f"{mes_num}_"
                        + re.sub(
                            r"[^A-Za-z0-9_-]",
                            "_",
                            grado_actual
                        )
                    )
                )

            st.markdown("---")
            st.markdown(
                f"### 📥 Descargar reportes "
                f"de {nombre_mes} {anio_actual}"
            )

            with st.container():
                excel_bytes_mes = (
                    generar_excel_bytes(
                        df_mes,
                        titulo_reporte=(
                            "REPORTE DE ASISTENCIA - "
                            f"{nombre_mes.upper()} "
                            f"{anio_actual}"
                        )
                    )
                )

                if excel_bytes_mes is not None:
                    st.download_button(
                        label=(
                            f"📊 Descargar Excel "
                            f"Profesional de "
                            f"{nombre_mes}"
                        ),
                        data=excel_bytes_mes,
                        file_name=(
                            f"asistencia_{anio_actual}_"
                            f"{mes_num:02d}_"
                            f"{nombre_mes}.xlsx"
                        ),
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.spreadsheetml.sheet"
                        ),
                        use_container_width=True,
                        type="primary",
                        key=(
                            f"btn_excel_profesional_"
                            f"{mes_num}"
                        ),
                    )



# ============================================================
# REPORTES V5.0
# ============================================================
def construir_resumen_asistencia(df, registros, clases):
    """Calcula asistencia por estudiante usando los registros existentes."""
    if df.empty:
        return pd.DataFrame()
    df = df.copy()
    df["codigo"] = df["codigo"].astype(str)
    conteo = df.groupby("codigo").size().to_dict()
    filas = []
    for r in registros:
        codigo = str(r.get("codigo", ""))
        clases_est = [c for c in clases if estudiante_en_clase(r, c)]
        total_clases = len(clases_est)
        presentes = conteo.get(codigo, 0)
        # Se usa el número de registros como indicador de asistencias realizadas.
        filas.append({
            "Código": codigo,
            "Estudiante": f"{r.get('nombres','')} {r.get('apellidos','')}".strip(),
            "Nivel": r.get("nivel", ""),
            "Grado": r.get("grado", ""),
            "Clases asignadas": total_clases,
            "Asistencias registradas": presentes,
        })
    resumen = pd.DataFrame(filas)
    if not resumen.empty:
        resumen["Porcentaje de registros"] = resumen.apply(
            lambda x: round((x["Asistencias registradas"] / x["Clases asignadas"]) * 100, 2)
            if x["Clases asignadas"] else 0, axis=1
        )
    return resumen


def panel_reportes_v5():
    st.markdown("### 📊 Reportes y estadísticas V5.0")
    st.caption("Consulta indicadores de asistencia y descarga un reporte de estudiantes.")
    clases = clases_permitidas_para_usuario()
    registros = registros_docente_actual() if st.session_state.get("rol") == "docente" else listar_registros()
    df = cargar_asistencias()
    if st.session_state.get("rol") == "docente":
        ids = {str(c.get("id")) for c in clases}
        if not df.empty:
            df = df[df["clase_id"].astype(str).isin(ids)].copy()
    resumen = construir_resumen_asistencia(df, registros, clases)
    if resumen.empty:
        st.info("Aún no hay datos suficientes para generar estadísticas.")
        return
    c1, c2, c3 = st.columns(3)
    c1.metric("👨‍🎓 Estudiantes", len(resumen))
    c2.metric("✅ Registros de asistencia", int(resumen["Asistencias registradas"].sum()))
    c3.metric("📚 Clases disponibles", len(clases))
    st.dataframe(resumen, use_container_width=True, hide_index=True)
    excel = generar_excel_bytes(
        df,
        titulo_reporte="REPORTE V5.0 DE ASISTENCIA"
    )
    if excel is not None:
        st.download_button(
            "📥 Descargar reporte Excel V5.0",
            data=excel,
            file_name="reporte_asistencia_v5.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

# ============================================================
# PORTAL DOCENTE
# ============================================================
def registros_docente_actual():
    docente = docente_actual()
    if not docente: return []
    clases = clases_permitidas_para_usuario()
    return [r for r in listar_registros() if any(estudiante_en_clase(r, c) for c in clases)]


def panel_inicio_docente():
    docente = docente_actual()
    if not docente: st.error("No se pudo cargar el perfil docente."); return
    estudiantes = registros_docente_actual(); clases = docente.get("clases", [])
    st.markdown(f"## 👋 Bienvenido, {html_lib.escape(docente.get('nombre', 'Docente'))}")
    st.caption("Portal docente · acceso limitado a las clases asignadas.")
    m1,m2,m3 = st.columns(3); m1.metric("👨‍🎓 Mis estudiantes", len(estudiantes)); m2.metric("📚 Mis clases", len(clases)); m3.metric("🎓 Grados", len(docente.get("grados", [])))
    st.markdown("---"); st.markdown("### 🚀 Accesos rápidos")
    b1,b2,b3 = st.columns(3)
    with b1:
        if st.button("👨‍🎓 Mis estudiantes", use_container_width=True): st.session_state.vista_actual="mis_estudiantes"; st.rerun()
    with b2:
        if st.button("📅 Registrar asistencia", use_container_width=True): st.session_state.vista_actual="asistencia"; st.rerun()
    with b3:
        if st.button("📚 Mis clases", use_container_width=True): st.session_state.vista_actual="mis_clases"; st.rerun()


def panel_docente_estudiantes():
    docente = docente_actual()
    if not docente:
        st.error("No se pudo cargar el perfil docente.")
        return
    clases = clases_permitidas_para_usuario()
    registros = registros_docente_actual()
    st.markdown("### 👨‍🎓 Mis estudiantes")
    st.caption("Use el botón 👁️ para consultar el carnet del estudiante.")
    if not registros:
        st.info("Todavía no tiene estudiantes registrados en sus clases.")
        return
    busqueda = st.text_input(
        "Buscar estudiante",
        placeholder="Nombre, código, grado o materia...",
        key="buscar_mis_estudiantes"
    )
    if busqueda:
        t = normalizar_texto(busqueda).casefold()
        registros = [
            r for r in registros
            if t in f"{r.get('codigo','')} {r.get('nombres','')} "
                 f"{r.get('apellidos','')} {r.get('grado','')} "
                 f"{r.get('materia','')}".casefold()
        ]
    for idx, c in enumerate(clases):
        grupo = [r for r in registros if estudiante_en_clase(r, c)]
        with st.expander(
            f"📖 {c.get('materia')} · {c.get('grado')} — "
            f"{len(grupo)} estudiante(s)",
            expanded=True
        ):
            if not grupo:
                st.info("No hay estudiantes en esta clase.")
                continue
            for j, r in enumerate(grupo):
                nombre = normalizar_texto(
                    f"{r.get('nombres','')} {r.get('apellidos','')}"
                )
                col_info, col_btn = st.columns([5, 1])
                with col_info:
                    st.markdown(
                        f"**{nombre}**  \\n"
                        f"Código: `{r.get('codigo','')}` · "
                        f"Grado: {r.get('grado','')}"
                    )
                with col_btn:
                    if st.button(
                        "👁️ Ver carnet",
                        key=f"ver_carnet_docente_{idx}_{j}_{r.get('codigo','')}",
                        use_container_width=True
                    ):
                        st.session_state["carnet_docente_ver"] = r.get("codigo")
            codigo_ver = st.session_state.get("carnet_docente_ver")
            if codigo_ver and any(r.get("codigo") == codigo_ver for r in grupo):
                datos = cargar_registro(codigo_ver)
                if datos:
                    st.markdown("---")
                    renderizar_carnet(datos)


def panel_docente_clases():
    docente = docente_actual()
    if not docente:
        st.error("No se pudo cargar el perfil docente.")
        return
    st.markdown("### 📚 Mis clases asignadas")
    clases = clases_permitidas_para_usuario()
    registros = registros_docente_actual()
    if not clases:
        st.info("No tiene clases asignadas.")
        return
    for idx, c in enumerate(clases):
        estudiantes = [r for r in registros if estudiante_en_clase(r, c)]
        with st.expander(
            f"📖 {c.get('materia')} · {c.get('grado')} — "
            f"{len(estudiantes)} estudiante(s)",
            expanded=True
        ):
            st.caption(f"Nivel: {c.get('nivel')} · Docente: {c.get('docente')}")
            for j, r in enumerate(estudiantes):
                a, b = st.columns([5, 1])
                with a:
                    st.write(
                        f"**{r.get('nombres','')} {r.get('apellidos','')}** "
                        f"· `{r.get('codigo','')}`"
                    )
                with b:
                    if st.button(
                        "👁️ Ver carnet",
                        key=f"ver_carnet_clase_{idx}_{j}_{r.get('codigo','')}",
                        use_container_width=True
                    ):
                        st.session_state["carnet_docente_ver"] = r.get("codigo")
            codigo_ver = st.session_state.get("carnet_docente_ver")
            if codigo_ver and any(r.get("codigo") == codigo_ver for r in estudiantes):
                datos = cargar_registro(codigo_ver)
                if datos:
                    st.markdown("---")
                    renderizar_carnet(datos)


# ============================================================
# MI CUENTA# ============================================================
# MI CUENTA
# ============================================================
def panel_mi_cuenta():
    st.markdown("### ⚙️ Mi cuenta")

    rol = st.session_state.get("rol")

    if rol == "docente":
        docente = docente_actual()
        if not docente:
            st.error("No se pudo cargar la cuenta docente.")
            return

        st.markdown(
            f"**Docente:** {html_lib.escape(docente.get('nombre', ''))}"
        )
        st.markdown(
            f"**Usuario:** {html_lib.escape(docente.get('usuario', ''))}"
        )
        st.caption(
            "Puede cambiar su contraseña. Los permisos de administrador "
            "no están disponibles para esta cuenta."
        )

        with st.form("form_cuenta_docente"):
            password_actual = st.text_input(
                "Contraseña actual",
                type="password"
            )
            nueva_password = st.text_input(
                "Nueva contraseña",
                type="password"
            )
            confirmar_password = st.text_input(
                "Confirmar nueva contraseña",
                type="password"
            )

            guardar = st.form_submit_button(
                "🔐 Cambiar contraseña",
                type="primary",
                use_container_width=True
            )

        if guardar:
            if not verificar_hash_password(
                password_actual,
                str(docente.get("password_hash", ""))
            ):
                st.error("La contraseña actual es incorrecta.")
            elif len(nueva_password) < 8:
                st.error("La nueva contraseña debe tener al menos 8 caracteres.")
            elif nueva_password != confirmar_password:
                st.error("Las contraseñas no coinciden.")
            else:
                docentes = cargar_docentes()
                for d in docentes:
                    if d.get("id") == docente.get("id"):
                        d["password_hash"] = hash_password(nueva_password)
                        d["requiere_cambio_password"] = False
                        break

                guardar_docentes(docentes)
                st.success("Contraseña actualizada correctamente.")
                st.rerun()

        return

    credenciales = cargar_credenciales_admin()

    st.markdown(
        f"**Usuario actual:** "
        f"{html_lib.escape(str(credenciales['usuario']))}"
    )

    if credenciales.get("requiere_cambio_password", False):
        st.warning(
            "Por seguridad, debe cambiar la contraseña inicial."
        )

    with st.form("form_cambiar_credenciales"):
        password_actual = st.text_input(
            "Contraseña actual",
            type="password"
        )
        nuevo_usuario = st.text_input(
            "Nuevo nombre de usuario",
            value=str(credenciales["usuario"])
        )
        nueva_password = st.text_input(
            "Nueva contraseña",
            type="password"
        )
        confirmar_password = st.text_input(
            "Confirmar nueva contraseña",
            type="password"
        )

        guardar = st.form_submit_button(
            "Guardar cambios",
            type="primary",
            use_container_width=True
        )

    if guardar:
        if not verificar_hash_password(
            password_actual,
            str(credenciales["password_hash"])
        ):
            st.error("La contraseña actual es incorrecta.")
        elif nueva_password and len(nueva_password) < 8:
            st.error("La nueva contraseña debe tener al menos 8 caracteres.")
        elif nueva_password and nueva_password != confirmar_password:
            st.error("Las contraseñas no coinciden.")
        elif not normalizar_texto(nuevo_usuario):
            st.error("El usuario no puede quedar vacío.")
        else:
            password_hash_final = (
                hash_password(nueva_password)
                if nueva_password
                else credenciales["password_hash"]
            )

            guardar_credenciales_admin(
                normalizar_texto(nuevo_usuario),
                password_hash_final,
                requiere_cambio_password=False
            )

            st.session_state.username = normalizar_texto(nuevo_usuario)
            st.session_state.password_inicial_temporal = None
            st.success("Datos actualizados correctamente.")
            st.rerun()


# ============================================================
# NAVEGACIÓN
# ============================================================
NAV_ITEMS_ADMIN = [
    ("inicio", "🏠", "Inicio", panel_inicio),
    ("registro", "📝", "Registrar carnet", pantalla_registro),
    ("docentes", "👩‍🏫", "Docentes", panel_docentes),
    ("datos", "📋", "Carnets registrados", panel_datos),
    ("asistencia", "📅", "Asistencia", panel_asistencia),
    ("historial", "🗓️", "Historial mensual", panel_historial_asistencia),
    ("reportes_v5", "📊", "Reportes V5.0", panel_reportes_v5),
    ("cuenta", "⚙️", "Mi cuenta", panel_mi_cuenta),
]

NAV_ITEMS_DOCENTE = [
    ("inicio_docente", "🏠", "Inicio", panel_inicio_docente),
    ("asistencia", "📅", "Registrar asistencia", panel_asistencia),
    ("mis_estudiantes", "👨‍🎓", "Mis estudiantes", panel_docente_estudiantes),
    ("mis_clases", "📚", "Mis clases", panel_docente_clases),
    ("reportes_v5", "📊", "Reportes V5.0", panel_reportes_v5),
    ("cuenta", "⚙️", "Mi cuenta", panel_mi_cuenta),
]



def renderizar_sidebar():
    items = (
        NAV_ITEMS_ADMIN
        if st.session_state.get("rol") == "admin"
        else NAV_ITEMS_DOCENTE
    )

    with st.sidebar:
        if LOGO_PATH:
            st.image(LOGO_PATH, width=56)

        st.markdown(
            f"""
            <div class='sidebar-nombre-colegio'>
                <div class='l1'>{NOMBRE_COLEGIO_LINEA1}</div>
                <div class='l2'>{NOMBRE_COLEGIO_LINEA2}</div>
                <div class='l3'>{NOMBRE_COLEGIO_LINEA3}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        rol = st.session_state.get("rol")
        etiqueta_rol = "Administrador" if rol == "admin" else "Docente"

        st.markdown(
            f"""
            <div style='padding:10px 12px;margin:8px 0 14px;
                        border-radius:12px;background:rgba(255,255,255,.08);'>
                <div style='font-size:10px;opacity:.7;'>CUENTA ACTIVA</div>
                <div style='font-weight:800;'>👤 {html_lib.escape(
                    str(st.session_state.get("username") or "")
                )}</div>
                <div style='font-size:11px;opacity:.75;'>{etiqueta_rol}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

        for clave, icono, etiqueta, _ in items:
            activo = st.session_state.vista_actual == clave

            if st.button(
                f"{icono}   {etiqueta}",
                key=f"nav_{rol}_{clave}",
                use_container_width=True,
                type="primary" if activo else "secondary"
            ):
                st.session_state.vista_actual = clave
                st.session_state.docente_editando = None
                st.session_state.carnet_editando = None
                st.rerun()

        st.markdown("<hr>", unsafe_allow_html=True)

        if st.button(
            "⏻   Cerrar sesión",
            key="btn_logout",
            use_container_width=True
        ):
            st.query_params.clear()
            st.session_state.logged_in = False
            st.session_state.rol = None
            st.session_state.username = None
            st.session_state.docente_id = None
            st.session_state.pagina = "login_admin"
            st.session_state.vista_actual = "inicio"
            st.session_state.ultimo_codigo_registrado = None
            st.session_state.docente_editando = None
            st.session_state.carnet_editando = None
            st.rerun()


# ============================================================
# FLUJO PRINCIPAL
# ============================================================
codigo_desde_qr = (
    st.query_params.get("codigo")
    or st.session_state.get("codigo_escaneado")
)

if (
    codigo_desde_qr
    and not st.session_state.logged_in
):
    pantalla_ver_carnet_publico(codigo_desde_qr)

elif (
    st.session_state.pagina == "login_admin"
    and not st.session_state.logged_in
):
    pantalla_login_admin()

elif (
    st.session_state.pagina == "app"
    and st.session_state.logged_in
):
    renderizar_sidebar()
    barra_superior()

    items = (
        NAV_ITEMS_ADMIN
        if st.session_state.get("rol") == "admin"
        else NAV_ITEMS_DOCENTE
    )

    funciones_por_vista = {
        clave: funcion
        for clave, _, _, funcion in items
    }

    vista = (
        st.session_state.vista_actual
        if st.session_state.vista_actual in funciones_por_vista
        else (
            "inicio"
            if st.session_state.get("rol") == "admin"
            else "inicio_docente"
        )
    )

    # Defensa adicional: un docente nunca puede invocar una vista administrativa.
    st.session_state.vista_actual = vista
    funciones_por_vista[vista]()

else:
    st.session_state.pagina = "login_admin"
    st.session_state.logged_in = False
    st.session_state.rol = None
    st.session_state.username = None
    st.session_state.docente_id = None
    st.rerun()
